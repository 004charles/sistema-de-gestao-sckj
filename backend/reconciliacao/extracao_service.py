"""Extração de texto de documentos em vários formatos (Fase 2).

Suporta PDF (pdfplumber), Excel (openpyxl), CSV (nativo) e XML (AGT).
Devolução uniforme: {'texto', 'formato', 'erro'}.
"""
import csv
import logging
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

from . import pdf_service

logger = logging.getLogger('reconciliacao')

FORMATOS = {
    '.pdf': 'pdf',
    '.xlsx': 'xlsx',
    '.csv': 'csv',
    '.xml': 'xml',
}

EXTENSOES_ACEITAS = tuple(FORMATOS.keys())

MAX_CELULAS_POR_LINHA = 60


def detectar_formato(nome_arquivo):
    nome = (nome_arquivo or '').lower()
    for extensao, formato in FORMATOS.items():
        if nome.endswith(extensao):
            return formato
    return None


def _extrair_pdf(caminho):
    texto = pdf_service.extrair_texto_pdf(caminho)
    return texto or ''


def _extrair_xlsx(caminho):
    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        linhas_txt = []
        for ws in wb.worksheets:
            linhas_txt.append(f'[Folha: {ws.title}]')
            for linha in ws.iter_rows(values_only=True):
                celulas = [
                    str(c).strip()
                    for c in linha[:MAX_CELULAS_POR_LINHA]
                    if c is not None and str(c).strip()
                ]
                if celulas:
                    linhas_txt.append(' | '.join(celulas))
        return '\n'.join(linhas_txt)
    finally:
        wb.close()


def _extrair_csv(caminho):
    for codificacao in ('utf-8-sig', 'cp1252', 'latin-1'):
        try:
            with open(caminho, 'r', encoding=codificacao, newline='') as f:
                amostra = f.read(4096)
                f.seek(0)
                try:
                    dialeto = csv.Sniffer().sniff(amostra, delimiters=';,\t|')
                except csv.Error:
                    dialeto = csv.excel
                linhas = []
                for linha in csv.reader(f, dialeto):
                    celulas = [c.strip() for c in linha if c and c.strip()]
                    if celulas:
                        linhas.append(' | '.join(celulas))
            return '\n'.join(linhas)
        except UnicodeDecodeError:
            continue
    return ''


def _nome_local(tag):
    return tag.split('}')[-1]


def _extrair_xml(caminho):
    raiz = ET.parse(caminho).getroot()
    valores = []
    nomes_tags = []
    for elem in raiz.iter():
        nomes_tags.append(_nome_local(elem.tag))
        texto = (elem.text or '').strip()
        if texto:
            valores.append(texto)
    partes = valores
    if not partes:
        partes = sorted(set(nomes_tags))
    else:
        partes = partes + [' '.join(sorted(set(nomes_tags)))]
    return '\n'.join(partes)


def extrair_conteudo(caminho, nome_arquivo):
    """Extrai texto do ficheiro. Nunca lança excepção: devolve {'texto','formato','erro'}."""
    formato = detectar_formato(nome_arquivo)
    if formato is None:
        return {'texto': '', 'formato': None, 'erro': 'Formato não suportado'}
    try:
        if formato == 'pdf':
            texto = _extrair_pdf(caminho)
        elif formato == 'xlsx':
            texto = _extrair_xlsx(caminho)
        elif formato == 'csv':
            texto = _extrair_csv(caminho)
        else:
            texto = _extrair_xml(caminho)
        return {'texto': texto or '', 'formato': formato, 'erro': None}
    except Exception as exc:
        logger.exception('Erro ao extrair %s', nome_arquivo)
        return {'texto': '', 'formato': formato, 'erro': f'Erro ao extrair texto: {exc}'}
