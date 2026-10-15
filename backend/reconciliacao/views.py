import logging
from decimal import Decimal
from django.conf import settings
from django.http import HttpResponse
from rest_framework import viewsets, status
from rest_framework.decorators import api_view, action, permission_classes
from rest_framework.response import Response
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.permissions import AllowAny, IsAuthenticated
from rest_framework.permissions import DjangoModelPermissions
from rest_framework.viewsets import ReadOnlyModelViewSet
from rest_framework.authtoken.models import Token
from rest_framework.exceptions import PermissionDenied
from django.contrib.auth.models import User, Group
from django.contrib.auth import authenticate
from django.db.models import Q
from .models import (
    Empresa, EmpresaUtilizador, Contabilidade, DeclaracaoAGT,
    Reconciliacao, ReconciliacaoDetalhe, MapeamentoConta,
    DocumentoUpload, AnaliseIA, Regra, Ocorrencia,
    RegistroAcessoAdmin, LancamentoContabilistico, LinhaLancamento,
    Factura
)
from .serializers import (
    EmpresaSerializer,
    ContabilidadeSerializer, DeclaracaoAGTSerializer,
    ReconciliacaoSerializer, ReconciliacaoDetalheSerializer,
    MapeamentoContaSerializer, LoginSerializer,
    ReconciliacaoExecutarSerializer,
    LancamentoContabilisticoSerializer, LinhaLancamentoSerializer,
    FacturaSerializer
)
from . import pdf_service
from . import extracao_service
from . import classificador_service
from . import cobertura_service
from . import validacao_documental
from . import dossie_service
from . import regra_engine
from . import calendario_service
from . import numeros_service
from . import taxas
from . import auditoria_trace
from . import ai_auditoria
from . import consolidacao_anual
from . import auto_cadastro_service
from . import factura_ai_service

logger = logging.getLogger('reconciliacao')


def registrar_tentativa_acesso(request, rota, resultado='NEGADO', motivo=''):
    """Regista no log e na base de dados as tentativas de acesso às áreas de administração."""
    try:
        user = request.user if getattr(request, 'user', None) and request.user.is_authenticated else None
        data = getattr(request, 'data', {}) or {}
        username = user.username if user else data.get('username', '') or 'Anónimo'
        ip = request.META.get('HTTP_X_FORWARDED_FOR', request.META.get('REMOTE_ADDR', ''))
        if ',' in ip:
            ip = ip.split(',')[0].strip()
        user_agent = request.META.get('HTTP_USER_AGENT', '')[:255]

        empresa_nome = ''
        if user:
            empresa_acesso = getattr(user, 'acessos_empresas', None)
            if empresa_acesso and empresa_acesso.exists():
                empresa_nome = ', '.join([e.empresa.nome for e in empresa_acesso.select_related('empresa').all()[:2]])

        if not empresa_nome and data.get('empresa_nome'):
            empresa_nome = str(data.get('empresa_nome'))[:255]

        RegistroAcessoAdmin.objects.create(
            utilizador=user,
            username=username,
            empresa_nome=empresa_nome,
            ip_origem=ip,
            user_agent=user_agent,
            rota_tentada=rota[:255],
            motivo=motivo[:255],
            resultado=resultado,
        )
        logger.warning(
            f"[AUDITORIA DE SEGURANÇA] Tentativa de acesso {resultado} à rota '{rota}' "
            f"por utilizador '{username}' (Empresa: '{empresa_nome}', IP: {ip}). Motivo: {motivo}"
        )
    except Exception as e:
        logger.error(f"Erro ao registar auditoria de acesso admin: {e}")


def empresas_acessiveis(user):
    if not user or not user.is_authenticated:
        return Empresa.objects.none()
    if user.is_superuser or user.is_staff or user.groups.filter(name__in=['Super Admin', 'Administrador']).exists():
        return Empresa.objects.all()
    # Se o utilizador é cliente de empresa, restringe estritamente às suas empresas vinculadas
    vinculos = Empresa.objects.filter(
        acessos__utilizador=user,
        acessos__ativo=True,
    ).distinct()
    if vinculos.exists():
        return vinculos
    return Empresa.objects.none()


def verificar_permissao(user, codename):
    if not user or not user.is_authenticated:
        return False
    if user.is_superuser:
        return True
    return user.has_perm(f'reconciliacao.{codename}')


def documentos_acessiveis(user):
    if not user or not user.is_authenticated:
        return DocumentoUpload.objects.none()
    if user.is_superuser:
        return DocumentoUpload.objects.all()
    return DocumentoUpload.objects.filter(empresa__in=empresas_acessiveis(user))


def _inteiro_opcional(valor):
    try:
        return int(valor)
    except (TypeError, ValueError):
        return None


@api_view(['POST'])
def login_view(request):
    serializer = LoginSerializer(data=request.data)
    if serializer.is_valid():
        username = serializer.validated_data['username']
        password = serializer.validated_data['password']
        user = authenticate(username=username, password=password)
        if user is not None and user.is_active:
            token, _ = Token.objects.get_or_create(user=user)
            grupos = list(user.groups.values_list('name', flat=True))
            is_admin = bool(
                user.is_superuser
                or user.is_staff
                or any(g in ['Super Admin', 'Administrador', 'Auditor Chefe'] for g in grupos)
            )
            empresas_ids = list(empresas_acessiveis(user).values_list('id', flat=True))
            return Response({
                'success': True,
                'user': {
                    'id': user.id,
                    'username': user.username,
                    'email': user.email,
                    'first_name': user.first_name,
                    'last_name': user.last_name,
                    'token': token.key,
                    'is_superuser': user.is_superuser,
                    'is_staff': user.is_staff,
                    'is_admin': is_admin,
                    'grupos': grupos,
                    'empresas_acesso': empresas_ids,
                }
            })
        return Response({'error': 'Credenciais inválidas'}, status=status.HTTP_401_UNAUTHORIZED)
    return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


login_view.cls.permission_classes = [AllowAny]


class EmpresaViewSet(viewsets.ModelViewSet):
    """Gestão de empresas: listagem, cadastro, edição e consulta."""

    queryset = Empresa.objects.all()
    serializer_class = EmpresaSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        queryset = empresas_acessiveis(self.request.user)
        nome = self.request.query_params.get('nome')
        nif = self.request.query_params.get('nif')
        estado = self.request.query_params.get('estado')
        if nome:
            queryset = queryset.filter(nome__icontains=nome)
        if nif:
            queryset = queryset.filter(nif__icontains=nif)
        if estado:
            queryset = queryset.filter(estado=estado)
        return queryset

    def perform_create(self, serializer):
        user = self.request.user
        is_admin = bool(
            user.is_superuser
            or user.is_staff
            or user.groups.filter(name__in=['Super Admin', 'Administrador']).exists()
        )
        if not is_admin and not user.has_perm('reconciliacao.add_empresa'):
            raise PermissionDenied('Apenas administradores podem cadastrar empresas.')
        empresa = serializer.save()
        if not user.is_superuser:
            EmpresaUtilizador.objects.get_or_create(
                empresa=empresa,
                utilizador=user,
                defaults={'perfil_acesso': 'ADMIN', 'ativo': True}
            )

    def perform_update(self, serializer):
        user = self.request.user
        is_admin = bool(
            user.is_superuser
            or user.is_staff
            or user.groups.filter(name__in=['Super Admin', 'Administrador']).exists()
        )
        if not is_admin and not user.has_perm('reconciliacao.change_empresa'):
            raise PermissionDenied('Apenas administradores podem alterar dados de empresas.')
        serializer.save()

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        user = request.user
        is_admin = bool(
            user.is_superuser
            or user.is_staff
            or user.groups.filter(name__in=['Super Admin', 'Administrador']).exists()
        )
        if not is_admin and not user.has_perm('reconciliacao.delete_empresa'):
            raise PermissionDenied('Apenas administradores podem eliminar empresas.')

        try:
            from django.db import transaction
            nome_empresa = instance.nome
            with transaction.atomic():
                # Remove documentos vinculados para desimpedir chave protegida
                instance.documentos.all().delete()
                instance.delete()
            logger.info(f"Empresa '{nome_empresa}' eliminada com sucesso por {user.username}")
            return Response(status=status.HTTP_204_NO_CONTENT)
        except Exception as e:
            logger.error(f"Erro ao eliminar empresa {instance.id}: {e}")
            return Response(
                {'error': f'Não foi possível eliminar a empresa: {str(e)}'},
                status=status.HTTP_400_BAD_REQUEST
            )


class ContabilidadeViewSet(viewsets.ModelViewSet):
    queryset = Contabilidade.objects.select_related('empresa').all()
    serializer_class = ContabilidadeSerializer
    permission_classes = [IsAuthenticated, DjangoModelPermissions]

    def get_queryset(self):
        queryset = Contabilidade.objects.select_related('empresa').filter(
            empresa__in=empresas_acessiveis(self.request.user)
        )
        empresa_id = self.request.query_params.get('empresa_id')
        ano = self.request.query_params.get('ano')
        mes = self.request.query_params.get('mes')
        if empresa_id:
            queryset = queryset.filter(empresa_id=empresa_id)
        if ano:
            queryset = queryset.filter(ano=ano)
        if mes:
            queryset = queryset.filter(mes=mes)
        return queryset


class DeclaracaoAGTViewSet(viewsets.ModelViewSet):
    queryset = DeclaracaoAGT.objects.select_related('empresa').all()
    serializer_class = DeclaracaoAGTSerializer
    permission_classes = [IsAuthenticated, DjangoModelPermissions]

    def get_queryset(self):
        queryset = DeclaracaoAGT.objects.select_related('empresa').filter(
            empresa__in=empresas_acessiveis(self.request.user)
        )
        empresa_id = self.request.query_params.get('empresa_id')
        ano = self.request.query_params.get('ano')
        mes = self.request.query_params.get('mes')
        if empresa_id:
            queryset = queryset.filter(empresa_id=empresa_id)
        if ano:
            queryset = queryset.filter(ano=ano)
        if mes:
            queryset = queryset.filter(mes=mes)
        return queryset


class ReconciliacaoViewSet(viewsets.ModelViewSet):
    queryset = Reconciliacao.objects.select_related('empresa').prefetch_related('detalhes').all()
    serializer_class = ReconciliacaoSerializer
    permission_classes = [IsAuthenticated, DjangoModelPermissions]

    def get_queryset(self):
        queryset = Reconciliacao.objects.select_related('empresa').prefetch_related('detalhes').filter(
            empresa__in=empresas_acessiveis(self.request.user)
        )
        empresa_id = self.request.query_params.get('empresa_id')
        ano = self.request.query_params.get('ano')
        mes = self.request.query_params.get('mes')
        status_filter = self.request.query_params.get('status')
        if empresa_id:
            queryset = queryset.filter(empresa_id=empresa_id)
        if ano:
            queryset = queryset.filter(ano=ano)
        if mes:
            queryset = queryset.filter(mes=mes)
        if status_filter:
            queryset = queryset.filter(status=status_filter)
        return queryset

    @action(detail=False, methods=['post'])
    def executar(self, request):
        serializer = ReconciliacaoExecutarSerializer(data=request.data)
        if serializer.is_valid():
            empresa_id = serializer.validated_data['empresa_id']
            ano = serializer.validated_data['ano']
            mes = serializer.validated_data['mes']
            resultado = self._executar_reconciliacao(empresa_id, ano, mes)
            return Response(resultado)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    def _executar_reconciliacao(self, empresa_id, ano, mes):
        res = dossie_service.executar_reconciliacao_automatica(empresa_id, ano, mes)
        if not res:
            return {'error': f'Dados insuficientes para reconciliação em {mes}/{ano}'}
        reconciliacao = Reconciliacao.objects.get(pk=res['id'])
        return ReconciliacaoSerializer(reconciliacao).data

    def _normalizar_contabilidade(self, dados):
        resultado = {
            'iva_liquidado': 0,
            'iva_dedutivel': 0,
            'iva_apuramento': 0,
            'iva_pagar': 0,
            'iva_recuperar': 0,
        }
        for item in dados:
            codigo = item.conta_codigo.upper()
            saldo = abs(float(item.saldo))
            if 'LIQUIDADO' in codigo or '24.4.1' in codigo:
                resultado['iva_liquidado'] += saldo
            elif 'DEDUTIVEL' in codigo or '24.4.2' in codigo:
                resultado['iva_dedutivel'] += saldo
            elif 'APURAMENTO' in codigo or '24.4.3' in codigo:
                resultado['iva_apuramento'] += saldo
            elif 'A PAGAR' in codigo or '24.4.4' in codigo:
                resultado['iva_pagar'] += saldo
            elif 'A RECUPERAR' in codigo or '24.4.5' in codigo:
                resultado['iva_recuperar'] += saldo
        return resultado


class ReconciliacaoDetalheViewSet(ReadOnlyModelViewSet):
    queryset = ReconciliacaoDetalhe.objects.all()
    serializer_class = ReconciliacaoDetalheSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        queryset = ReconciliacaoDetalhe.objects.filter(
            reconciliacao__empresa__in=empresas_acessiveis(self.request.user)
        )
        reconciliacao_id = self.request.query_params.get('reconciliacao_id')
        if reconciliacao_id:
            queryset = queryset.filter(reconciliacao_id=reconciliacao_id)
        return queryset


class MapeamentoContaViewSet(viewsets.ModelViewSet):
    queryset = MapeamentoConta.objects.all()
    serializer_class = MapeamentoContaSerializer
    permission_classes = [IsAuthenticated, DjangoModelPermissions]


@api_view(['GET'])
def dashboard_view(request, empresa_id):
    try:
        empresa = empresas_acessiveis(request.user).get(pk=empresa_id)
    except Empresa.DoesNotExist:
        return Response({'error': 'Empresa não encontrada'}, status=404)

    total_reconciliacoes = Reconciliacao.objects.filter(empresa_id=empresa_id).count()
    reconciliacoes_ok = Reconciliacao.objects.filter(empresa_id=empresa_id, status='CONCILIADO').count()
    reconciliacoes_divergencia = Reconciliacao.objects.filter(empresa_id=empresa_id, status='DIVERGENCIA').count()

    ultimas_reconciliacoes = Reconciliacao.objects.filter(empresa_id=empresa_id)[:5]

    return Response({
        'empresa': EmpresaSerializer(empresa).data,
        'estatisticas': {
            'total_reconciliacoes': total_reconciliacoes,
            'conciliadas': reconciliacoes_ok,
            'divergencias': reconciliacoes_divergencia,
        },
        'ultimas_reconciliacoes': ReconciliacaoSerializer(ultimas_reconciliacoes, many=True).data
    })


@api_view(['GET'])
def relatorio_resumo_view(request):
    empresas = empresas_acessiveis(request.user)
    dados = []
    for empresa in empresas:
        total = Reconciliacao.objects.filter(empresa=empresa).count()
        ok = Reconciliacao.objects.filter(empresa=empresa, status='CONCILIADO').count()
        div = Reconciliacao.objects.filter(empresa=empresa, status='DIVERGENCIA').count()
        dados.append({
            'empresa': EmpresaSerializer(empresa).data,
            'total': total,
            'conciliadas': ok,
            'divergencias': div,
        })
    return Response(dados)


TIPOS_UNICOS_POR_PERIODO = ('BALANCETE', 'MODELO7', 'IMPOSTO_INDUSTRIAL', 'DECLARACAO_IRT')


def _documento_duplicado(empresa, ano, mes, tipo, nome_arquivo=None, excluir_pk=None):
    """Documento do mesmo tipo no mesmo período que bloqueia novo upload.

    Documentos ilegíveis ou com erro não bloqueiam: permitem tentar de novo.
    Para declarações periódicas e balancetes (BALANCETE, MODELO7, etc.), admite apenas 1 por período.
    Para documentos de suporte e arquivo (faturas, comprovativos, folhas, etc.), permite múltiplos,
    barrando apenas se já existir ficheiro com o mesmo nome no mesmo período.
    """
    qs = DocumentoUpload.objects.filter(
        empresa=empresa, ano=ano, mes=mes, tipo=tipo,
    ).exclude(estado__in=('ERRO', 'ILEGIVEL'))
    if excluir_pk:
        qs = qs.exclude(pk=excluir_pk)
    if tipo in TIPOS_UNICOS_POR_PERIODO:
        return qs.order_by('-id').first()
    if nome_arquivo:
        return qs.filter(nome_arquivo=nome_arquivo).order_by('-id').first()
    return None


def _resposta_duplicado(duplicado, tipo, ano, mes):
    return Response(
        {
            'error': (
                f'Já existe um documento do tipo {tipo} para {mes:02d}/{ano} '
                f'("{duplicado.nome_arquivo}"). Elimine-o antes de subir outro.'
            )
        },
        status=status.HTTP_409_CONFLICT,
    )


@api_view(['POST'])
def upload_documento_view(request):
    if not verificar_permissao(request.user, 'add_documentoupload'):
        return Response(
            {'error': 'Sem permissão para carregar documentos.'},
            status=status.HTTP_403_FORBIDDEN
        )
    arquivo = request.FILES.get('arquivo')
    tipo = (request.data.get('tipo') or '').strip()
    empresa_id = request.data.get('empresa_id') or request.data.get('empresa')
    ano = _inteiro_opcional(request.data.get('ano'))
    mes = _inteiro_opcional(request.data.get('mes'))
    substituir = str(request.data.get('substituir', '')).lower() in ('1', 'true', 'yes', 'on')

    if not arquivo:
        return Response(
            {'error': 'arquivo é obrigatório'},
            status=status.HTTP_400_BAD_REQUEST
        )

    tipos_validos = [c[0] for c in DocumentoUpload.TIPO_CHOICES]
    if tipo and tipo not in tipos_validos:
        return Response(
            {'error': f'tipo deve ser um de: {", ".join(tipos_validos)}'},
            status=status.HTTP_400_BAD_REQUEST
        )

    empresa = None
    empresa_existe_inacessivel = False
    if empresa_id and str(empresa_id).strip().lower() not in ('auto', '0', 'none', '', 'null', 'undefined'):
        try:
            empresa = empresas_acessiveis(request.user).filter(pk=empresa_id).first()
            if not empresa and Empresa.objects.filter(pk=empresa_id).exists():
                empresa_existe_inacessivel = True
        except Exception:
            empresa = None

    if empresa_existe_inacessivel:
        return Response(
            {'error': 'Empresa não encontrada ou sem acesso.'},
            status=status.HTTP_404_NOT_FOUND
        )

    empresa_recem_criada = False
    mensagem_cadastro = ''

    # Se a empresa não foi especificada OU a empresa informada já não existe (ex: foi eliminada):
    # inspeciona o documento para identificar e pré-cadastrar automaticamente!
    if empresa is None:
        import tempfile
        import os
        extensao = ('.' + arquivo.name.rsplit('.', 1)[-1].lower() if '.' in arquivo.name else '')
        with tempfile.NamedTemporaryFile(suffix=extensao, delete=False) as tmp:
            for chunk in arquivo.chunks():
                tmp.write(chunk)
            tmp_caminho = tmp.name
        try:
            conteudo = extracao_service.extrair_conteudo(tmp_caminho, arquivo.name)
            texto_up = conteudo.get('texto') or ''
            dados_cad = auto_cadastro_service.extrair_dados_empresa_do_texto(texto_up, arquivo.name)
            if dados_cad.get('nif') or dados_cad.get('nome'):
                emp_auto, criada, msg = auto_cadastro_service.obter_ou_cadastrar_empresa_automatica(
                    dados_cad, user=request.user, auto_criar=True
                )
                if emp_auto:
                    empresa = emp_auto
                    empresa_id = emp_auto.id
                    empresa_recem_criada = criada
                    mensagem_cadastro = msg
            if ano is None or mes is None:
                ano_det, mes_det, _ = dossie_service.detectar_periodo_e_nif(texto_up)
                ano = ano or ano_det
                mes = mes or mes_det
        finally:
            if os.path.exists(tmp_caminho):
                os.remove(tmp_caminho)

    if not empresa:
        if not empresa_id or str(empresa_id).strip().lower() in ('auto', '0', 'none', '', 'null', 'undefined'):
            return Response(
                {
                    'error': (
                        'empresa_id é obrigatório. '
                        'Não foi possível identificar automaticamente os dados da empresa (NIF/Nome) '
                        'no documento para efetuar o pré-cadastro.'
                    )
                },
                status=status.HTTP_400_BAD_REQUEST
            )
        return Response(
            {
                'error': (
                    'Empresa não encontrada ou sem acesso. '
                    'Não foi possível identificar automaticamente os dados da empresa (NIF/Nome) '
                    'no documento para efetuar o pré-cadastro.'
                )
            },
            status=status.HTTP_404_NOT_FOUND
        )

    if ano is None or mes is None or not (2000 <= ano <= 2100) or not (1 <= mes <= 12):
        return Response(
            {'error': 'ano (2000-2100) e mes (1-12) são obrigatórios'},
            status=status.HTTP_400_BAD_REQUEST
        )

    if tipo:
        nome_doc = request.data.get('nome') or arquivo.name
        duplicado = _documento_duplicado(empresa, ano, mes, tipo, nome_arquivo=nome_doc)
        if duplicado is not None:
            if substituir:
                duplicado.delete()
            else:
                return _resposta_duplicado(duplicado, tipo, ano, mes)

    extensao = ('.' + arquivo.name.rsplit('.', 1)[-1].lower()
                if '.' in arquivo.name else '')
    if extensao not in extracao_service.EXTENSOES_ACEITAS:
        aceites = ', '.join(extracao_service.EXTENSOES_ACEITAS)
        return Response(
            {'error': f'Formato não suportado ({extensao or "sem extensão"}). Aceites: {aceites}'},
            status=status.HTTP_400_BAD_REQUEST
        )

    documento = DocumentoUpload.objects.create(
        arquivo=arquivo,
        empresa=empresa,
        ano=ano,
        mes=mes,
        tipo=tipo or 'OUTRO',
        estado='PROCESSANDO',
        nome_arquivo=request.data.get('nome') or arquivo.name,
    )

    resultado = extracao_service.extrair_conteudo(documento.arquivo.path, arquivo.name)
    texto = resultado['texto']

    sugestao = classificador_service.classificar(documento.nome_arquivo, texto)
    documento.tipo_sugerido = sugestao['tipo']
    if not tipo:
        documento.tipo = sugestao['tipo']
        duplicado = _documento_duplicado(
            empresa, ano, mes, documento.tipo, nome_arquivo=documento.nome_arquivo, excluir_pk=documento.pk
        )
        if duplicado is not None:
            if substituir:
                duplicado.delete()
            else:
                tipo_efetivo = documento.tipo
                arquivo_salvo = documento.arquivo
                documento.delete()
                arquivo_salvo.delete(save=False)
                return _resposta_duplicado(duplicado, tipo_efetivo, ano, mes)

    if resultado['erro'] or (texto or '').startswith('Erro ao extrair texto'):
        documento.estado = 'ERRO'
    elif not (texto or '').strip():
        documento.estado = 'ILEGIVEL'
    else:
        documento.estado = 'VALIDADO'

    # Processamento especial de Dossiê Mensal Completo (Envelope Fiscal multi-documentos)
    is_dossie = (tipo == 'DOSSIE_MENSAL') or (
        extensao == '.pdf' and dossie_service.is_dossie_mensal(documento.arquivo.path)
    )
    if is_dossie and documento.estado == 'VALIDADO':
        documento.tipo = 'DOSSIE_MENSAL'
        documento.tipo_sugerido = 'DOSSIE_MENSAL'
        documento.texto_extraido = texto
        documento.save()
        res_dossie = dossie_service.desmembrar_e_processar_dossie(
            documento.arquivo.path, empresa, ano, mes
        )
        return Response({
            'success': True,
            'is_dossie': True,
            'mensagem': res_dossie['mensagem'],
            'ano': res_dossie.get('ano') or documento.ano,
            'mes': res_dossie.get('mes') or documento.mes,
            'empresa_id': empresa.id,
            'documento': {
                'id': documento.id,
                'tipo': 'DOSSIE_MENSAL',
                'tipo_sugerido': 'DOSSIE_MENSAL',
                'classificado_automaticamente': not bool(tipo),
                'formato': resultado['formato'],
                'estado': 'VALIDADO',
                'empresa': empresa.id,
                'empresa_nome': empresa.nome,
                'ano': res_dossie.get('ano') or documento.ano,
                'mes': res_dossie.get('mes') or documento.mes,
                'nome_arquivo': documento.nome_arquivo,
            },
            'documentos_criados': res_dossie['documentos_criados'],
            'auditoria': res_dossie.get('auditoria'),
            'reconciliacao': res_dossie.get('reconciliacao'),
            'analise_ia_id': res_dossie.get('analise_ia_id'),
        })

    # Validação estrita de conteúdo para BALANCETE e MODELO7 em documentos validados
    if documento.estado == 'VALIDADO' and documento.tipo in ('BALANCETE', 'MODELO7'):
        valido, motivo_erro = validacao_documental.validar_documento_para_tipo(
            documento.tipo, texto, documento.nome_arquivo
        )
        if not valido:
            arquivo_salvo = documento.arquivo
            documento.delete()
            arquivo_salvo.delete(save=False)
            logger.warning(f"Upload rejeitado por validação estrita ({documento.tipo}): {motivo_erro}")
            return Response(
                {'error': motivo_erro},
                status=status.HTTP_400_BAD_REQUEST
            )

    documento.texto_extraido = texto
    documento.save()

    # Automação de Ponta a Ponta: Se o documento for validado e contiver ano e mês,
    # sincroniza Contabilidade/Declarações, executa o motor de 15 regras fiscais,
    # reconcilia e processa IA se ambos os documentos já estiverem disponíveis.
    auditoria_auto = None
    reconciliacao_auto = None
    analise_ia_id = None
    if documento.estado == 'VALIDADO' and documento.ano and documento.mes:
        try:
            doc_b = documento if documento.tipo == 'BALANCETE' else DocumentoUpload.objects.filter(
                empresa=empresa, ano=documento.ano, mes=documento.mes, tipo='BALANCETE'
            ).order_by('-id').first()
            doc_m = documento if documento.tipo in ('MODELO7', 'COMPROVATIVOS') else DocumentoUpload.objects.filter(
                empresa=empresa, ano=documento.ano, tipo__in=['MODELO7', 'COMPROVATIVOS']
            ).order_by('-id').first()

            _sincronizar_dados_documentos(empresa, documento.ano, documento.mes, doc_b, doc_m)
            auditoria_auto = regra_engine.executar(empresa, documento.ano, documento.mes)
            reconciliacao_auto = dossie_service.executar_reconciliacao_automatica(empresa, documento.ano, documento.mes)

            if doc_b and doc_m and doc_b.texto_extraido and doc_m.texto_extraido:
                if not AnaliseIA.objects.filter(documento_contabilidade=doc_b, documento_agt=doc_m, status='CONCLUIDA').exclude(resultado='').exists():
                    res_ia = pdf_service.analisar_documentos(doc_b.texto_extraido, doc_m.texto_extraido, 'pt')
                    if res_ia.get('success'):
                        analise_ia_obj = AnaliseIA.objects.create(
                            documento_contabilidade=doc_b,
                            documento_agt=doc_m,
                            resultado=res_ia.get('analise'),
                            status='CONCLUIDA'
                        )
                        analise_ia_id = analise_ia_obj.id
        except Exception as e:
            logger.warning(f"Erro no auto-processamento de auditoria pós-upload: {e}")

    return Response({
        'success': True,
        'documento': {
            'id': documento.id,
            'tipo': documento.tipo,
            'tipo_sugerido': documento.tipo_sugerido,
            'classificado_automaticamente': not bool(tipo),
            'formato': resultado['formato'],
            'estado': documento.estado,
            'empresa': empresa.id,
            'empresa_nome': empresa.nome,
            'empresa_recem_criada': empresa_recem_criada,
            'mensagem_cadastro': mensagem_cadastro,
            'ano': documento.ano,
            'mes': documento.mes,
            'nome_arquivo': documento.nome_arquivo,
            'texto_preview': texto[:500] + '...' if len(texto) > 500 else texto,
        },
        'auditoria': auditoria_auto,
        'reconciliacao': reconciliacao_auto,
        'analise_ia_id': analise_ia_id,
    })


@api_view(['POST'])
def detectar_documento_view(request):
    """Inspeciona o documento antes do upload final para identificar automaticamente:
    - Empresa (através do NIF ou Razão Social no texto)
    - Período (Ano e Mês)
    - Tipo de documento (Dossiê Completo, Balancete, Modelo 7, etc.)
    """
    arquivo = request.FILES.get('arquivo')
    if not arquivo:
        return Response({'error': 'Ficheiro não fornecido'}, status=status.HTTP_400_BAD_REQUEST)

    import tempfile
    import os
    import unicodedata
    extensao = ('.' + arquivo.name.rsplit('.', 1)[-1].lower() if '.' in arquivo.name else '')
    
    with tempfile.NamedTemporaryFile(suffix=extensao, delete=False) as tmp:
        for chunk in arquivo.chunks():
            tmp.write(chunk)
        tmp_path = tmp.name

    try:
        resultado = extracao_service.extrair_conteudo(tmp_path, arquivo.name)
        texto = resultado.get('texto') or ''

        is_dossie = False
        if extensao == '.pdf':
            is_dossie = dossie_service.is_dossie_mensal(tmp_path)

        if is_dossie:
            tipo = 'DOSSIE_MENSAL'
            tipo_nome = 'Dossiê Mensal Completo (Envelope Fiscal Único)'
        else:
            sugestao = classificador_service.classificar(arquivo.name, texto)
            tipo = sugestao['tipo']
            tipo_nome = dict(DocumentoUpload.TIPO_CHOICES).get(tipo, tipo)

        # Detectar Empresa
        empresas = empresas_acessiveis(request.user)
        empresa_detetada = None
        texto_norm = unicodedata.normalize('NFKD', texto).lower()

        for emp in empresas:
            if emp.nif and emp.nif in texto:
                empresa_detetada = emp
                break

        if not empresa_detetada:
            for emp in empresas:
                nome_simplificado = ''.join(c for c in unicodedata.normalize('NFKD', emp.nome) if c.isalnum() or c.isspace()).lower()
                partes = [p for p in nome_simplificado.split() if len(p) > 3][:3]
                if partes and all(p in texto_norm for p in partes):
                    empresa_detetada = emp
                    break

        empresa_recem_criada = False
        mensagem_cadastro = ''

        # Auto-cadastro inteligente: se a empresa não existe na base, extrai dados e cadastra automaticamente!
        if not empresa_detetada:
            dados_cad = auto_cadastro_service.extrair_dados_empresa_do_texto(texto, arquivo.name)
            if dados_cad.get('nif') or dados_cad.get('nome'):
                emp_auto, criada, msg = auto_cadastro_service.obter_ou_cadastrar_empresa_automatica(
                    dados_cad, user=request.user, auto_criar=True
                )
                if emp_auto:
                    empresa_detetada = emp_auto
                    empresa_recem_criada = criada
                    mensagem_cadastro = msg

        ano, mes, _ = dossie_service.detectar_periodo_e_nif(texto)

        return Response({
            'success': True,
            'arquivo_nome': arquivo.name,
            'is_dossie': is_dossie,
            'tipo': tipo,
            'tipo_nome': tipo_nome,
            'empresa_id': empresa_detetada.id if empresa_detetada else None,
            'empresa_nome': empresa_detetada.nome if empresa_detetada else None,
            'empresa_nif': empresa_detetada.nif if empresa_detetada else None,
            'empresa_recem_criada': empresa_recem_criada,
            'mensagem_cadastro': mensagem_cadastro,
            'ano': ano,
            'mes': mes,
        })
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def _sincronizar_dados_documentos(empresa, ano, mes, doc_balancete=None, doc_agt=None):
    """Sincroniza automaticamente Contabilidade e DeclaracaoAGT com os ficheiros PDF carregados."""
    try:
        if doc_balancete and doc_balancete.texto_extraido and not Contabilidade.objects.filter(empresa=empresa, ano=ano, mes=mes).exists():
            b_dados = pdf_service.extrair_dados_balancete(doc_balancete.texto_extraido)
            contas = [
                ('6211', 'Prestações de Serviço - Mercado Nacional', Decimal('0'), b_dados['vendas'], -b_dados['vendas']),
                ('4511', 'Caixa Kwanza', Decimal('0'), b_dados['caixa_saldo_credor'], -b_dados['caixa_saldo_credor']),
                ('3453', 'IVA Liquidado - Operações Gerais', b_dados['iva_liquidado'], b_dados['iva_liquidado'], Decimal('0')),
                ('3451', 'IVA Suportado / Dedutível', b_dados['iva_suportado'], b_dados['iva_suportado'], Decimal('0')),
                ('3456', 'IVA a Pagar de Apuramento', Decimal('0'), b_dados['iva_pagar'], -b_dados['iva_pagar']),
                ('3431', 'Imposto de Rendimento do Trabalho (IRT)', Decimal('0'), b_dados['irt_retido'], -b_dados['irt_retido']),
                ('3492', 'Segurança Social (INSS)', Decimal('0'), b_dados['inss_retido'], -b_dados['inss_retido']),
                ('3493', 'Retenção na Fonte Prestadores (6,5%)', Decimal('0'), b_dados['retencao_prestadores'], -b_dados['retencao_prestadores']),
                ('3413', 'Retenção na Fonte Clientes', b_dados['retencao_clientes'], Decimal('0'), b_dados['retencao_clientes']),
                ('75234', 'Honorários e Avenças (Prestadores)', b_dados['honorarios_prestadores'], Decimal('0'), b_dados['honorarios_prestadores']),
                ('722', 'Remunerações - Pessoal', b_dados['remuneracoes'], Decimal('0'), b_dados['remuneracoes']),
                ('75312', 'IVA Lançado em Custos', b_dados['iva_custo'], Decimal('0'), b_dados['iva_custo']),
            ]
            for cod, desc, deb, cred, sal in contas:
                if deb > 0 or cred > 0:
                    Contabilidade.objects.create(
                        empresa=empresa,
                        ano=ano,
                        mes=mes,
                        conta_codigo=cod,
                        conta_descricao=desc,
                        valor_debito=deb,
                        valor_credito=cred,
                        saldo=sal,
                    )
        if doc_agt and doc_agt.texto_extraido and not DeclaracaoAGT.objects.filter(empresa=empresa, ano=ano, mes=mes).exists():
            m_dados = pdf_service.extrair_dados_modelo7(doc_agt.texto_extraido)
            if m_dados['base_tributavel'] > 0 or m_dados['iva_liquidado'] > 0 or m_dados['iva_pagar'] > 0:
                DeclaracaoAGT.objects.create(
                    empresa=empresa,
                    ano=ano,
                    mes=mes,
                    nif=empresa.nif,
                    razao_social=empresa.nome,
                    regime_iva=empresa.regime_iva or 'Regime Geral',
                    iva_liquidado=m_dados['iva_liquidado'],
                    iva_dedutivel=m_dados['iva_dedutivel'],
                    iva_apurado=m_dados['iva_liquidado'] - m_dados['iva_dedutivel'],
                    iva_pagar=m_dados['iva_pagar'],
                    iva_recuperar=m_dados['iva_recuperar'],
                )
    except Exception as e:
        logger.warning(f"Erro na sincronização de documentos para contabilidade/agt: {e}")


@api_view(['POST'])
def analisar_documentos_view(request):
    if not verificar_permissao(request.user, 'add_analiseia'):
        return Response(
            {'error': 'Sem permissão para criar análises.'},
            status=status.HTTP_403_FORBIDDEN
        )
    logger.info("=== INÍCIO DA ANÁLISE DE DOCUMENTOS ===")
    
    doc_contab_id = request.data.get('documento_contabilidade_id')
    doc_agt_id = request.data.get('documento_agt_id')
    idioma = request.data.get('idioma', 'pt')

    logger.info(f"Recebido: doc_contab_id={doc_contab_id}, doc_agt_id={doc_agt_id}, idioma={idioma}")

    if not doc_contab_id or not doc_agt_id:
        logger.error("IDs dos documentos não fornecidos")
        return Response(
            {'error': 'documento_contabilidade_id e documento_agt_id são obrigatórios'},
            status=status.HTTP_400_BAD_REQUEST
        )

    try:
        docs_usuario = documentos_acessiveis(request.user)
        doc_contab = docs_usuario.filter(pk=doc_contab_id).first()
        doc_agt = docs_usuario.filter(pk=doc_agt_id).first()
        if doc_contab is None or doc_agt is None:
            raise DocumentoUpload.DoesNotExist
        logger.info(f"Documentos encontrados: contab={doc_contab.nome_arquivo}, agt={doc_agt.nome_arquivo}")
    except DocumentoUpload.DoesNotExist:
        logger.error(f"Documento não encontrado: contab_id={doc_contab_id}, agt_id={doc_agt_id}")
        return Response(
            {'error': 'Documentos não encontrados ou sem acesso.'},
            status=status.HTTP_404_NOT_FOUND
        )

    # Re-extrair texto se vazio
    if not (doc_contab.texto_extraido or '').strip() and doc_contab.arquivo:
        try:
            res_c = extracao_service.extrair_conteudo(doc_contab.arquivo.path, doc_contab.nome_arquivo)
            if res_c.get('texto'):
                doc_contab.texto_extraido = res_c['texto']
                doc_contab.save(update_fields=['texto_extraido'])
        except Exception as e:
            logger.warning(f"Erro ao extrair contab: {e}")

    if not (doc_agt.texto_extraido or '').strip() and doc_agt.arquivo:
        try:
            res_a = extracao_service.extrair_conteudo(doc_agt.arquivo.path, doc_agt.nome_arquivo)
            if res_a.get('texto'):
                doc_agt.texto_extraido = res_a['texto']
                doc_agt.save(update_fields=['texto_extraido'])
        except Exception as e:
            logger.warning(f"Erro ao extrair agt: {e}")

    logger.info(f"Texto extraído contabilidade: {len(doc_contab.texto_extraido or '')} chars")
    logger.info(f"Texto extraído AGT: {len(doc_agt.texto_extraido or '')} chars")

    # Validação estrita antes de acionar a Inteligência Artificial
    val_c, err_c = validacao_documental.validar_balancete(doc_contab.texto_extraido, doc_contab.nome_arquivo)
    if not val_c:
        logger.warning(f"Análise rejeitada: Balancete inválido ({err_c})")
        return Response(
            {'error': f"Documento de contabilidade inválido: {err_c}"},
            status=status.HTTP_400_BAD_REQUEST
        )

    val_a, err_a = validacao_documental.validar_modelo7(doc_agt.texto_extraido, doc_agt.nome_arquivo)
    if not val_a:
        logger.warning(f"Análise rejeitada: Modelo 7 inválido ({err_a})")
        return Response(
            {'error': f"Documento da AGT inválido: {err_a}"},
            status=status.HTTP_400_BAD_REQUEST
        )

    analise = AnaliseIA.objects.create(
        documento_contabilidade=doc_contab,
        documento_agt=doc_agt,
        status='PENDENTE'
    )
    logger.info(f"Análise criada com ID: {analise.id}")

    logger.info("Chamando pdf_service.analisar_documentos...")
    resultado = pdf_service.analisar_documentos(
        doc_contab.texto_extraido,
        doc_agt.texto_extraido,
        idioma
    )

    logger.info(f"Resultado da análise: success={resultado.get('success')}")

    if resultado.get('success'):
        analise.resultado = resultado['analise']
        analise.status = 'CONCLUIDA'
        analise.save()
        DocumentoUpload.objects.filter(
            pk__in=[doc_contab.pk, doc_agt.pk]
        ).update(estado='UTILIZADO')
        # Sincronizar dados estruturados com tabelas contabilidade e agt
        emp = doc_contab.empresa or doc_agt.empresa
        ano_doc = doc_contab.ano or doc_agt.ano
        mes_doc = doc_contab.mes or doc_agt.mes
        if emp and ano_doc and mes_doc:
            _sincronizar_dados_documentos(emp, ano_doc, mes_doc, doc_contab, doc_agt)
        logger.info("Análise concluída com sucesso")
        return Response({
            'success': True,
            'analise_id': analise.id,
            'resultado': resultado['analise'],
            'model': resultado.get('model'),
        })
    else:
        analise.status = 'ERRO'
        analise.save()
        error_msg = resultado.get('error')
        logger.error(f"Erro na análise: {error_msg}")
        return Response(
            {'error': error_msg},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR
        )


@api_view(['POST'])
def analisar_documento_unico_view(request):
    if not verificar_permissao(request.user, 'add_analiseia'):
        return Response(
            {'error': 'Sem permissão para criar análises.'},
            status=status.HTTP_403_FORBIDDEN
        )
    documento_id = request.data.get('documento_id')
    idioma = request.data.get('idioma', 'pt')

    if not documento_id:
        return Response(
            {'error': 'documento_id é obrigatório'},
            status=status.HTTP_400_BAD_REQUEST
        )

    try:
        documento = documentos_acessiveis(request.user).filter(pk=documento_id).first()
        if documento is None:
            raise DocumentoUpload.DoesNotExist
    except DocumentoUpload.DoesNotExist:
        return Response(
            {'error': 'Documento não encontrado'},
            status=status.HTTP_404_NOT_FOUND
        )

    resultado = pdf_service.analisar_documento_unico(
        documento.texto_extraido,
        documento.tipo,
        idioma
    )

    if resultado.get('success'):
        return Response({
            'success': True,
            'resultado': resultado['analise'],
            'model': resultado.get('model'),
        })
    else:
        return Response(
            {'error': resultado.get('error')},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR
        )


@api_view(['GET'])
def listar_documentos_view(request):
    documentos = documentos_acessiveis(request.user)

    empresa_id = _inteiro_opcional(request.query_params.get('empresa'))
    if empresa_id:
        documentos = documentos.filter(empresa_id=empresa_id)
    ano = _inteiro_opcional(request.query_params.get('ano'))
    if ano:
        documentos = documentos.filter(ano=ano)
    mes = _inteiro_opcional(request.query_params.get('mes'))
    if mes:
        documentos = documentos.filter(mes=mes)
    tipo = request.query_params.get('tipo')
    if tipo:
        documentos = documentos.filter(tipo=tipo)
    estado = request.query_params.get('estado')
    if estado:
        documentos = documentos.filter(estado=estado)

    dados = []
    for doc in documentos:
        dados.append({
            'id': doc.id,
            'tipo': doc.tipo,
            'tipo_sugerido': doc.tipo_sugerido,
            'estado': doc.estado,
            'empresa': doc.empresa_id,
            'empresa_nome': doc.empresa.nome if doc.empresa else None,
            'ano': doc.ano,
            'mes': doc.mes,
            'nome_arquivo': doc.nome_arquivo,
            'arquivo_url': doc.arquivo.url if doc.arquivo else None,
            'created_at': doc.created_at,
        })
    return Response(dados)


@api_view(['DELETE'])
def deletar_documento_view(request, documento_id):
    if not verificar_permissao(request.user, 'delete_documentoupload'):
        return Response(
            {'error': 'Sem permissão para eliminar documentos.'},
            status=status.HTTP_403_FORBIDDEN
        )
    documento = documentos_acessiveis(request.user).filter(pk=documento_id).first()
    if documento is None:
        return Response({'error': 'Documento não encontrado'}, status=status.HTTP_404_NOT_FOUND)
    empresa_id, ano, mes = documento.empresa_id, documento.ano, documento.mes
    documento.delete()

    # Sem documentos no período não podem sobrar resultados de análise.
    limpos = {'ocorrencias': 0, 'reconciliacoes': 0}
    if ano is not None and mes is not None:
        restam = DocumentoUpload.objects.filter(
            empresa_id=empresa_id, ano=ano, mes=mes
        ).exists()
        if not restam:
            limpos['ocorrencias'], _ = Ocorrencia.objects.filter(
                empresa_id=empresa_id, ano=ano, mes=mes
            ).delete()
            limpos['reconciliacoes'], _ = Reconciliacao.objects.filter(
                empresa_id=empresa_id, ano=ano, mes=mes
            ).delete()
    return Response({'success': True, 'resultados_limpos': limpos})


@api_view(['GET'])
def historico_analises_view(request):
    docs_usuario = documentos_acessiveis(request.user)
    analises = AnaliseIA.objects.select_related(
        'documento_contabilidade', 'documento_agt'
    ).filter(
        Q(documento_contabilidade__in=docs_usuario)
        | Q(documento_agt__in=docs_usuario)
    )
    dados = []
    for a in analises:
        dados.append({
            'id': a.id,
            'documento_contabilidade': {
                'id': a.documento_contabilidade.id if a.documento_contabilidade else None,
                'nome_arquivo': a.documento_contabilidade.nome_arquivo if a.documento_contabilidade else None,
            },
            'documento_agt': {
                'id': a.documento_agt.id if a.documento_agt else None,
                'nome_arquivo': a.documento_agt.nome_arquivo if a.documento_agt else None,
            },
            'status': a.status,
            'resultado': a.resultado[:500] + '...' if a.resultado and len(a.resultado) > 500 else a.resultado,
            'created_at': a.created_at,
        })
    return Response(dados)


@api_view(['DELETE'])
def deletar_analise_view(request, analise_id):
    if not verificar_permissao(request.user, 'delete_analiseia'):
        return Response(
            {'error': 'Sem permissão para eliminar análises.'},
            status=status.HTTP_403_FORBIDDEN
        )
    try:
        docs_usuario = documentos_acessiveis(request.user)
        analise = AnaliseIA.objects.filter(pk=analise_id).filter(
            Q(documento_contabilidade__in=docs_usuario)
            | Q(documento_agt__in=docs_usuario)
        ).first()
        if analise is None:
            return Response({'error': 'Análise não encontrada'}, status=404)
        analise.delete()
        return Response({'success': True})
    except Exception:
        return Response({'error': 'Análise não encontrada'}, status=404)


@api_view(['GET'])
def cobertura_view(request, empresa_id, ano, mes):
    """Cobertura de auditoria do período — Fase 3 (§10, §11)."""
    if not (2000 <= ano <= 2100) or not (1 <= mes <= 12):
        return Response(
            {'error': 'ano (2000-2100) e mes (1-12) inválidos'},
            status=status.HTTP_400_BAD_REQUEST
        )
    empresa = empresas_acessiveis(request.user).filter(pk=empresa_id).first()
    if empresa is None:
        return Response(
            {'error': 'Empresa não encontrada ou sem acesso.'},
            status=status.HTTP_404_NOT_FOUND
        )
    return Response(cobertura_service.calcular(empresa, ano, mes))


@api_view(['GET'])
def taxas_view(request):
    """Tabela de taxas e regimes em vigor (§12).

    Alimenta Definições → Parâmetros Fiscais, que até aqui tinha os valores
    escritos à mão no JSX e por isso envelhecia sem ninguém dar conta.
    """
    return Response(taxas.referencia())


@api_view(['GET'])
def numeros_view(request, empresa_id, ano, mes):
    """Os números do período: em risco, erros, OK, prazos (§33).

    É a resposta única que alimenta a página de Auditoria. Substitui o hábito
    anterior de três páginas pedirem o mesmo dashboard e cada uma repintar os
    mesmos dados de maneira diferente.
    """
    if not (2000 <= ano <= 2100 and 1 <= mes <= 12):
        return Response(
            {'error': 'ano (2000-2100) e mes (1-12) inválidos'},
            status=status.HTTP_400_BAD_REQUEST
        )
    empresa = empresas_acessiveis(request.user).filter(pk=empresa_id).first()
    if empresa is None:
        return Response(
            {'error': 'Empresa não encontrada ou sem acesso.'},
            status=status.HTTP_404_NOT_FOUND
        )
    return Response(numeros_service.calcular(empresa, ano, mes))


@api_view(['GET'])
def prazos_view(request, empresa_id, ano, mes):
    """Calendário fiscal do período: o que falta entregar e quando vence."""
    if not (2000 <= ano <= 2100 and 1 <= mes <= 12):
        return Response(
            {'error': 'ano (2000-2100) e mes (1-12) inválidos'},
            status=status.HTTP_400_BAD_REQUEST
        )
    empresa = empresas_acessiveis(request.user).filter(pk=empresa_id).first()
    if empresa is None:
        return Response(
            {'error': 'Empresa não encontrada ou sem acesso.'},
            status=status.HTTP_404_NOT_FOUND
        )
    return Response(calendario_service.resumo(empresa, ano, mes))


@api_view(['GET'])
def dashboard_auditoria_periodo_view(request, empresa_id, ano, mes):
    """Retorna a visão unificada do dashboard de auditoria do período (§33)."""
    if not (2000 <= ano <= 2100 and 1 <= mes <= 12):
        return Response(
            {'error': 'ano (2000-2100) e mes (1-12) inválidos'},
            status=status.HTTP_400_BAD_REQUEST
        )
    empresa = empresas_acessiveis(request.user).filter(pk=empresa_id).first()
    if empresa is None:
        return Response(
            {'error': 'Empresa não encontrada ou sem acesso.'},
            status=status.HTTP_404_NOT_FOUND
        )

    # 1. Cobertura documental
    cobertura = cobertura_service.calcular(empresa, ano, mes)

    # 2. Ocorrências / Alertas do período
    ocorrencias_qs = Ocorrencia.objects.filter(
        empresa=empresa, ano=ano, mes=mes
    ).select_related('regra', 'regra__motor', 'regra__base_legal', 'empresa').order_by('-severidade', 'id')

    alertas_resumo = {
        'alto': ocorrencias_qs.filter(severidade='ALTO').count(),
        'medio': ocorrencias_qs.filter(severidade='MEDIO').count(),
        'baixo': ocorrencias_qs.filter(severidade='BAIXO').count(),
        'informativo': ocorrencias_qs.filter(severidade='INFORMATIVO').count(),
        'total': ocorrencias_qs.count(),
    }

    # 3. Status semafórico por motor (§33)
    motores_status = []
    for m in cobertura['motores']:
        occs_motor = ocorrencias_qs.filter(regra__motor__codigo=m['codigo'])
        m_copy = dict(m)
        m_copy['alertas_alto'] = occs_motor.filter(severidade='ALTO').count()
        m_copy['alertas_medio'] = occs_motor.filter(severidade='MEDIO').count()
        m_copy['alertas_total'] = occs_motor.count()

        if m['estado'] in ('VAZIO', 'SEM_MATRIZ'):
            m_copy['status_semaforo'] = 'NAO_ANALISADO'  # ⚪
        elif m_copy['alertas_alto'] > 0:
            m_copy['status_semaforo'] = 'ALERTA'         # 🔴
        elif m_copy['alertas_medio'] > 0 or m['estado'] == 'PARCIAL':
            m_copy['status_semaforo'] = 'ATENCAO'        # 🟡
        else:
            m_copy['status_semaforo'] = 'CONFORME'       # 🟢
        motores_status.append(m_copy)

    # 4. Dois documentos mestres do período (Contabilidade e Portal AGT)
    doc_balancete = DocumentoUpload.objects.filter(
        empresa=empresa, ano=ano, mes=mes, tipo='BALANCETE'
    ).order_by('-id').first()
    if not doc_balancete:
        doc_balancete = DocumentoUpload.objects.filter(
            empresa=empresa, ano=ano, mes=mes, tipo='DOSSIE_MENSAL'
        ).order_by('-id').first()

    # Procurar AnaliseIA existente para a contabilidade deste período
    analise_ia_obj = None
    if doc_balancete:
        analise_ia_obj = AnaliseIA.objects.filter(
            documento_contabilidade=doc_balancete,
            status='CONCLUIDA'
        ).exclude(resultado='').order_by('-id').first()

    doc_agt = DocumentoUpload.objects.filter(
        empresa=empresa, ano=ano, mes=mes, tipo__in=['MODELO7', 'COMPROVATIVOS']
    ).order_by('-id').first()

    # Se não houver Modelo 7 no mês exato mas há AnaliseIA, resolve o doc_agt a partir da análise
    if not doc_agt and analise_ia_obj:
        doc_agt = analise_ia_obj.documento_agt

    # Se ainda não houver, procura Modelo 7 mais recente da empresa no mesmo ano (ex: Junho anexado em Julho)
    if not doc_agt:
        doc_agt = DocumentoUpload.objects.filter(
            empresa=empresa, ano=ano, tipo__in=['MODELO7', 'COMPROVATIVOS']
        ).order_by('-mes', '-id').first()

    # Se ainda não houver, procura Dossiê Mensal do período
    if not doc_agt:
        doc_agt = DocumentoUpload.objects.filter(
            empresa=empresa, ano=ano, mes=mes, tipo='DOSSIE_MENSAL'
        ).order_by('-id').first()

    if not analise_ia_obj and doc_balancete and doc_agt:
        analise_ia_obj = AnaliseIA.objects.filter(
            documento_contabilidade=doc_balancete,
            documento_agt=doc_agt,
            status='CONCLUIDA'
        ).exclude(resultado='').order_by('-id').first()

    # Sincronizar dados estruturados se ainda não gravados
    _sincronizar_dados_documentos(empresa, ano, mes, doc_balancete, doc_agt)

    # 5. Indicadores contábeis e fiscais do período
    from . import regra_engine
    contexto, _ = regra_engine.construir_contexto(empresa, ano, mes)

    documentos_base = {
        'balancete': {
            'presente': doc_balancete is not None,
            'id': doc_balancete.id if doc_balancete else None,
            'nome': doc_balancete.nome_arquivo if doc_balancete else None,
            'estado': doc_balancete.estado if doc_balancete else None,
            'data': doc_balancete.created_at.strftime('%d/%m/%Y %H:%M') if doc_balancete and doc_balancete.created_at else None,
        },
        'portal_agt': {
            'presente': doc_agt is not None,
            'id': doc_agt.id if doc_agt else None,
            'nome': doc_agt.nome_arquivo if doc_agt else None,
            'tipo': doc_agt.tipo if doc_agt else None,
            'estado': doc_agt.estado if doc_agt else None,
            'data': doc_agt.created_at.strftime('%d/%m/%Y %H:%M') if doc_agt and doc_agt.created_at else None,
        },
    }

    return Response({
        'empresa': EmpresaSerializer(empresa).data,
        'ano': ano,
        'mes': mes,
        'cobertura': cobertura,
        'documentos_base': documentos_base,
        'analise_ia': analise_ia_obj.resultado if analise_ia_obj else None,
        'alertas_resumo': alertas_resumo,
        'motores': motores_status,
        'ocorrencias': [_serializar_ocorrencia(o, completo=True) for o in ocorrencias_qs],
        'indicadores': {
            'volume_negocios': str(contexto.get('volume_negocios_contab', '0.00')),
            'iva_liquidado': str(contexto.get('iva_liquidado', '0.00')),
            'iva_dedutivel': str(contexto.get('iva_dedutivel', '0.00')),
            'iva_pagar': str(contexto.get('iva_pagar', '0.00')),
            'caixa_saldo_credor': str(contexto.get('caixa_saldo_credor', '0.00')),
            'caixa_credor_anormal': bool(contexto.get('caixa_credor_anormal', 0)),
            'irt_contabilizado': str(contexto.get('irt_contabilizado', '0.00')),
            'inss_contabilizado': str(contexto.get('inss_contabilizado', '0.00')),
            'honorarios_prestadores': str(contexto.get('honorarios_prestadores', '0.00')),
            'retencao_prestadores': str(contexto.get('retencao_prestadores_contab', '0.00')),
            'remuneracoes_total': str(contexto.get('remuneracoes_total', '0.00')),
            'bancos_saldo': str(contexto.get('bancos_saldo', '0.00')),
            'exposicao_total_potencial': str(contexto.get('exposicao_total_potencial', '0.00')),
            'iva_suportado_contab': str(contexto.get('iva_suportado_contab', '0.00')),
            'retencoes_clientes_contab': str(contexto.get('retencoes_clientes_contab', '0.00')),
        }
    })


@api_view(['GET'])
def rasto_auditoria_view(request, empresa_id, ano, mes):
    """Rasto determinístico da auditoria técnica passo-a-passo (§32, §33)."""
    if not (2000 <= ano <= 2100 and 1 <= mes <= 12):
        return Response(
            {'error': 'ano (2000-2100) e mes (1-12) inválidos'},
            status=status.HTTP_400_BAD_REQUEST
        )
    empresa = empresas_acessiveis(request.user).filter(pk=empresa_id).first()
    if empresa is None:
        return Response(
            {'error': 'Empresa não encontrada ou sem acesso.'},
            status=status.HTTP_404_NOT_FOUND
        )
    numeros = numeros_service.calcular(empresa, ano, mes)
    rasto = auditoria_trace.gerar(empresa, ano, mes, numeros)
    return Response(rasto)


@api_view(['GET', 'POST'])
def parecer_ia_view(request, empresa_id, ano, mes):
    """Parecer e análise de riscos com IA (ou fallback determinístico) para contabilistas."""
    if not (2000 <= ano <= 2100 and 1 <= mes <= 12):
        return Response(
            {'error': 'ano (2000-2100) e mes (1-12) inválidos'},
            status=status.HTTP_400_BAD_REQUEST
        )
    empresa = empresas_acessiveis(request.user).filter(pk=empresa_id).first()
    if empresa is None:
        return Response(
            {'error': 'Empresa não encontrada ou sem acesso.'},
            status=status.HTTP_404_NOT_FOUND
        )
    numeros = numeros_service.calcular(empresa, ano, mes)
    rasto = auditoria_trace.gerar(empresa, ano, mes, numeros)
    parecer_resultado = ai_auditoria.parecer(rasto, numeros)
    riscos_resultado = ai_auditoria.riscos(rasto, numeros)

    return Response({
        'parecer': parecer_resultado,
        'riscos': riscos_resultado,
        'rasto': rasto,
        'numeros': numeros,
    })


@api_view(['GET'])
def admin_estatisticas_view(request):
    """Métricas consolidadas para a Área Administrativa."""
    user = request.user
    is_admin = bool(
        user.is_authenticated
        and (
            user.is_superuser
            or user.is_staff
            or user.groups.filter(name__in=['Super Admin', 'Administrador']).exists()
        )
    )
    if not is_admin:
        registrar_tentativa_acesso(
            request,
            rota='/api/admin/estatisticas/',
            resultado='NEGADO',
            motivo='Tentativa de aceder às métricas administrativas globais por utilizador de empresa.'
        )
        raise PermissionDenied('Acesso reservado ao Administrador do Sistema. As empresas não têm permissão para aceder à Gestão Global.')

    empresas_qs = Empresa.objects.all().order_by('nome')
    total_empresas = empresas_qs.count()
    ativas = empresas_qs.filter(estado='ATIVA').count()
    inativas = total_empresas - ativas
    total_users = User.objects.count()

    lista = []
    total_risco_global = Decimal('0.00')

    for emp in empresas_qs:
        docs_count = emp.documentos.count()
        ultimo_doc = emp.documentos.order_by('-ano', '-mes', '-created_at').first()
        ano_ref = ultimo_doc.ano if ultimo_doc else 2026
        mes_ref = ultimo_doc.mes if ultimo_doc else 7

        occs = Ocorrencia.objects.filter(empresa=emp, ano=ano_ref, mes=mes_ref).exclude(estado='IGNORADA')
        erros_count = occs.filter(severidade__in=['ALTO', 'MEDIO']).count()

        risco_empresa = Decimal('0.00')
        for occ in occs.filter(severidade__in=['ALTO', 'MEDIO']):
            try:
                risco_empresa += Decimal(str(occ.diferenca or 0))
            except Exception:
                pass

        total_risco_global += risco_empresa

        lista.append({
            'id': emp.id,
            'nome': emp.nome,
            'nif': emp.nif,
            'regime_iva': emp.regime_iva,
            'regime_nome': emp.caracteristicas_regime.get('nome', emp.regime_iva),
            'estado': emp.estado,
            'volume_negocios_anual': str(emp.volume_negocios_anual or '0.00'),
            'telefone': emp.telefone or '',
            'email': emp.email or '',
            'endereco': emp.endereco or '',
            'provincia': emp.provincia or '',
            'contabilista_nome': emp.contabilista_nome or '',
            'total_documentos': docs_count,
            'ultimo_periodo': f'{mes_ref:02d}/{ano_ref}' if docs_count > 0 else 'Sem dados',
            'erros_contagem': erros_count,
            'em_risco': str(risco_empresa),
        })

    return Response({
        'total_empresas': total_empresas,
        'empresas_ativas': ativas,
        'empresas_inativas': inativas,
        'total_utilizadores': total_users,
        'total_risco_global': str(total_risco_global),
        'empresas': lista,
    })


@api_view(['GET', 'POST'])
def admin_utilizadores_view(request):
    """Gestão de utilizadores e atribuição a empresas."""
    user = request.user
    is_admin = bool(
        user.is_authenticated
        and (
            user.is_superuser
            or user.is_staff
            or user.groups.filter(name__in=['Super Admin', 'Administrador']).exists()
        )
    )
    if not is_admin:
        registrar_tentativa_acesso(
            request,
            rota='/api/admin/utilizadores/',
            resultado='NEGADO',
            motivo='Tentativa de aceder à gestão de utilizadores globais.'
        )
        raise PermissionDenied('Acesso reservado ao Administrador do Sistema. As empresas não têm permissão para aceder à Gestão Global.')

    if request.method == 'GET':
        users = User.objects.all().order_by('-date_joined')
        data = []
        for u in users:
            vinculos = EmpresaUtilizador.objects.filter(utilizador=u, ativo=True).select_related('empresa')
            data.append({
                'id': u.id,
                'username': u.username,
                'email': u.email,
                'first_name': u.first_name,
                'last_name': u.last_name,
                'is_active': u.is_active,
                'is_staff': u.is_staff,
                'is_superuser': u.is_superuser,
                'grupos': list(u.groups.values_list('name', flat=True)),
                'empresas': [
                    {'id': v.empresa.id, 'nome': v.empresa.nome, 'nif': v.empresa.nif, 'perfil': v.perfil_acesso}
                    for v in vinculos
                ]
            })
        return Response(data)

    elif request.method == 'POST':
        username = request.data.get('username')
        password = request.data.get('password')
        email = request.data.get('email', '')
        first_name = request.data.get('first_name', '')
        last_name = request.data.get('last_name', '')
        papel = request.data.get('papel', 'Contabilista')
        empresas_ids = request.data.get('empresas_ids', [])

        if not username or not password:
            return Response({'error': 'Nome de utilizador e palavra-passe são obrigatórios.'}, status=400)
        if User.objects.filter(username=username).exists():
            return Response({'error': 'Já existe um utilizador registado com este nome.'}, status=400)

        novo_user = User.objects.create_user(
            username=username,
            password=password,
            email=email,
            first_name=first_name,
            last_name=last_name
        )

        grupo = Group.objects.filter(name=papel).first()
        if grupo:
            novo_user.groups.add(grupo)

        if papel in ['Super Admin', 'Administrador']:
            novo_user.is_staff = True
            novo_user.save()

        for emp_id in empresas_ids:
            emp = Empresa.objects.filter(pk=emp_id).first()
            if emp:
                EmpresaUtilizador.objects.create(
                    empresa=emp,
                    utilizador=novo_user,
                    perfil_acesso='ADMIN' if papel in ['Super Admin', 'Administrador'] else 'CONTABILISTA',
                    ativo=True
                )

        return Response({
            'success': True,
            'id': novo_user.id,
            'username': novo_user.username
        }, status=201)


@api_view(['PUT', 'PATCH', 'DELETE'])
def admin_utilizador_detalhe_view(request, user_id):
    """Editar ou eliminar utilizador (apenas admin)."""
    user = request.user
    if not (user.is_superuser or user.is_staff or user.groups.filter(name__in=['Super Admin', 'Administrador']).exists()):
        registrar_tentativa_acesso(
            request,
            rota=f'/api/admin/utilizadores/{user_id}/',
            resultado='NEGADO',
            motivo='Tentativa de editar ou eliminar utilizador sem privilégios master.'
        )
        raise PermissionDenied('Apenas o Administrador pode gerir utilizadores.')
    alvo = User.objects.filter(pk=user_id).first()
    if not alvo:
        return Response({'error': 'Utilizador não encontrado.'}, status=404)

    if request.method == 'DELETE':
        if alvo.id == user.id:
            return Response({'error': 'Não é possível eliminar o próprio utilizador com sessão iniciada.'}, status=400)
        alvo.delete()
        return Response({'success': True})

    # PUT / PATCH - Atualização do utilizador
    data = request.data
    if 'first_name' in data:
        alvo.first_name = data.get('first_name', '').strip()
    if 'last_name' in data:
        alvo.last_name = data.get('last_name', '').strip()
    if 'email' in data:
        alvo.email = data.get('email', '').strip()
    if 'is_active' in data:
        alvo.is_active = bool(data.get('is_active'))
    if 'password' in data and data.get('password'):
        alvo.set_password(data.get('password'))

    papel = data.get('papel')
    if papel:
        alvo.groups.clear()
        grupo = Group.objects.filter(name=papel).first()
        if grupo:
            alvo.groups.add(grupo)
        if papel in ['Super Admin', 'Administrador']:
            alvo.is_staff = True
            if papel == 'Super Admin':
                alvo.is_superuser = True
        else:
            if not user.is_superuser:
                alvo.is_staff = False
                alvo.is_superuser = False

    alvo.save()

    # Atualizar empresas vinculadas se fornecido
    if 'empresas_ids' in data:
        novas_empresas_ids = [int(i) for i in data.get('empresas_ids', []) if str(i).isdigit()]
        # Remove vínculos antigos
        EmpresaUtilizador.objects.filter(utilizador=alvo).exclude(empresa_id__in=novas_empresas_ids).delete()
        for emp_id in novas_empresas_ids:
            emp = Empresa.objects.filter(pk=emp_id).first()
            if emp:
                EmpresaUtilizador.objects.get_or_create(
                    empresa=emp,
                    utilizador=alvo,
                    defaults={
                        'perfil_acesso': 'ADMIN' if papel in ['Super Admin', 'Administrador'] else 'CONTABILISTA',
                        'ativo': True
                    }
                )

    return Response({
        'success': True,
        'id': alvo.id,
        'username': alvo.username,
        'email': alvo.email,
        'first_name': alvo.first_name,
        'last_name': alvo.last_name,
        'is_active': alvo.is_active
    })



def _serializar_regra(regra, completo=False):
    base = regra.base_legal
    dados = {
        'id': regra.id,
        'codigo': regra.codigo,
        'nome': regra.nome,
        'imposto': regra.imposto,
        'motor': regra.motor.codigo,
        'motor_nome': regra.motor.nome,
        'severidade': regra.severidade,
        'estado_validacao': regra.estado_validacao,
        'activa': regra.activa,
        'versao': regra.versao,
        'base_legal': {
            'legislacao': base.legislacao,
            'artigo': base.artigo,
            'texto': base.texto,
            'fonte_oficial': base.fonte_oficial,
            'validado_juridicamente': base.validado_juridicamente,
        },
    }
    if completo:
        dados.update({
            'descricao': regra.descricao,
            'condicao': regra.condicao,
            'formula': regra.formula,
            'recomendacao': regra.recomendacao,
            'validado_por': regra.validado_por.username if regra.validado_por else None,
            'validado_em': regra.validado_em,
            'validacao_nota': regra.validacao_nota,
        })
    return dados


def _serializar_ocorrencia(ocorrencia, completo=False):
    dados = {
        'id': ocorrencia.id,
        'empresa': ocorrencia.empresa_id,
        'empresa_nome': ocorrencia.empresa.nome,
        'ano': ocorrencia.ano,
        'mes': ocorrencia.mes,
        'regra': _serializar_regra(ocorrencia.regra, completo=completo),
        'severidade': ocorrencia.severidade,
        'valor_envolvido': str(ocorrencia.valor_envolvido),
        'estado': ocorrencia.estado,
        'versao_regra': ocorrencia.versao_regra,
        'created_at': ocorrencia.created_at,
        'revisado_por': (
            ocorrencia.revisado_por.username if ocorrencia.revisado_por else None
        ),
        'revisado_em': ocorrencia.revisado_em,
    }
    if completo:
        dados.update({
            'evidencia': ocorrencia.evidencia,
            'recomendacao': ocorrencia.recomendacao,
        })
    return dados


@api_view(['GET'])
def listar_ocorrencias_view(request):
    """Alertas do período (§21) — sempre dentro de empresas acessíveis."""
    ocorrencias = Ocorrencia.objects.filter(
        empresa__in=empresas_acessiveis(request.user)
    ).select_related('regra', 'regra__motor', 'regra__base_legal', 'empresa')

    empresa_id = _inteiro_opcional(request.query_params.get('empresa'))
    if empresa_id:
        ocorrencias = ocorrencias.filter(empresa_id=empresa_id)
    ano = _inteiro_opcional(request.query_params.get('ano'))
    if ano:
        ocorrencias = ocorrencias.filter(ano=ano)
    mes = _inteiro_opcional(request.query_params.get('mes'))
    if mes:
        ocorrencias = ocorrencias.filter(mes=mes)
    severidade = request.query_params.get('severidade')
    if severidade:
        ocorrencias = ocorrencias.filter(severidade=severidade)
    estado = request.query_params.get('estado')
    if estado:
        ocorrencias = ocorrencias.filter(estado=estado)

    return Response([
        _serializar_ocorrencia(o) for o in ocorrencias
    ])


@api_view(['GET', 'PATCH'])
def detalhe_ocorrencia_view(request, ocorrencia_id):
    ocorrencia = Ocorrencia.objects.filter(
        pk=ocorrencia_id,
        empresa__in=empresas_acessiveis(request.user),
    ).select_related('regra', 'regra__motor', 'regra__base_legal', 'empresa').first()
    if ocorrencia is None:
        return Response({'error': 'Ocorrência não encontrada'}, status=404)

    if request.method == 'GET':
        return Response(_serializar_ocorrencia(ocorrencia, completo=True))

    if not verificar_permissao(request.user, 'change_ocorrencia'):
        return Response(
            {'error': 'Sem permissão para rever ocorrências.'},
            status=status.HTTP_403_FORBIDDEN
        )
    novo_estado = request.data.get('estado')
    estados = [c[0] for c in Ocorrencia.ESTADO_CHOICES]
    if novo_estado not in estados:
        return Response(
            {'error': f'estado deve ser um de: {", ".join(estados)}'},
            status=status.HTTP_400_BAD_REQUEST
        )
    from django.utils import timezone
    ocorrencia.estado = novo_estado
    ocorrencia.revisado_por = request.user
    ocorrencia.revisado_em = timezone.now()
    ocorrencia.save()
    return Response(_serializar_ocorrencia(ocorrencia, completo=True))


@api_view(['POST'])
def executar_ocorrencias_view(request):
    """Executa o Rule Engine sobre (empresa, ano, mes) — só regras VALIDADAS."""
    if not verificar_permissao(request.user, 'add_ocorrencia'):
        return Response(
            {'error': 'Sem permissão para executar o motor de regras.'},
            status=status.HTTP_403_FORBIDDEN
        )
    empresa_id = _inteiro_opcional(request.data.get('empresa') or request.data.get('empresa_id'))
    ano = _inteiro_opcional(request.data.get('ano'))
    mes = _inteiro_opcional(request.data.get('mes'))
    if not empresa_id or ano is None or mes is None:
        return Response(
            {'error': 'empresa, ano e mes são obrigatórios'},
            status=status.HTTP_400_BAD_REQUEST
        )
    if not (2000 <= ano <= 2100 and 1 <= mes <= 12):
        return Response(
            {'error': 'ano (2000-2100) e mes (1-12) inválidos'},
            status=status.HTTP_400_BAD_REQUEST
        )
    empresa = empresas_acessiveis(request.user).filter(pk=empresa_id).first()
    if empresa is None:
        return Response(
            {'error': 'Empresa não encontrada ou sem acesso.'},
            status=status.HTTP_404_NOT_FOUND
        )
    from . import regra_engine
    resumo = regra_engine.executar(empresa, ano, mes)
    return Response({'success': True, 'empresa': empresa.id, 'ano': ano, 'mes': mes, **resumo})


@api_view(['GET'])
def listar_regras_view(request):
    regras = Regra.objects.select_related('motor', 'base_legal')
    estado = request.query_params.get('estado')
    if estado:
        regras = regras.filter(estado_validacao=estado)
    motor = request.query_params.get('motor')
    if motor:
        regras = regras.filter(motor__codigo=motor)
    imposto = request.query_params.get('imposto')
    if imposto:
        regras = regras.filter(imposto=imposto)
    return Response([
        _serializar_regra(r, completo=True) for r in regras
    ])


@api_view(['PATCH'])
def validar_regra_view(request, regra_id):
    """Gestão e validação de regras pelo administrador."""
    if not verificar_permissao(request.user, 'change_regra'):
        return Response(
            {'error': 'Sem permissão para alterar regras.'},
            status=status.HTTP_403_FORBIDDEN
        )
    regra = Regra.objects.filter(pk=regra_id).first()
    if regra is None:
        return Response({'error': 'Regra não encontrada'}, status=404)

    if 'activa' in request.data:
        regra.activa = bool(request.data.get('activa'))

    if 'severidade' in request.data:
        sev = request.data.get('severidade')
        if sev in [c[0] for c in Regra.SEVERIDADE_CHOICES]:
            regra.severidade = sev

    novo_estado = request.data.get('estado_validacao')
    if novo_estado:
        estados = [c[0] for c in Regra.ESTADO_VALIDACAO_CHOICES]
        if novo_estado not in estados:
            return Response(
                {'error': f'estado_validacao deve ser um de: {", ".join(estados)}'},
                status=status.HTTP_400_BAD_REQUEST
            )
        from django.utils import timezone
        regra.estado_validacao = novo_estado
        regra.validado_por = request.user
        regra.validado_em = timezone.now()
        if 'nota' in request.data:
            regra.validacao_nota = str(request.data.get('nota') or '')[:2000]

    regra.save()
    return Response(_serializar_regra(regra, completo=True))


CONFIG_FILE_PATH = settings.BASE_DIR / 'configuracoes_sistema.json'

DEFAULT_CONFIGURACOES = {
    'nome_sistema': 'SCKJ - Sistema de Reconciliação Contabilística & Fiscal',
    'versao_sistema': 'v2.4.0 (OGE Angola 2026)',
    'modo_manutencao': False,
    'idioma_padrao': 'pt',
    'fuso_horario': 'Africa/Luanda (GMT+1)',
    'limite_upload_mb': 50,
    'taxa_iva_geral': 14.0,
    'taxa_retencao_fonte_ii': 6.5,
    'taxa_inss_patronal': 8.0,
    'taxa_inss_trabalhador': 3.0,
    'limite_alerta_risco_kz': 5000000.0,
    'prazo_inss_dia': 10,
    'prazo_iva_dia': 31,
    'notificacoes_email_ativas': True,
    'email_alertas_fiscal': 'auditoria@sckj.co.ao',
}


@api_view(['GET', 'POST', 'PUT', 'PATCH'])
@permission_classes([IsAuthenticated])
def admin_configuracoes_view(request):
    """
    Endpoint para consulta e atualização das configurações globais do sistema administrativo.
    Apenas acessível a Super Administradores e utilizadores Staff.
    """
    user = request.user
    if not (user.is_superuser or user.is_staff or getattr(user, 'is_admin', False)):
        registrar_tentativa_acesso(
            request,
            rota='/api/admin/configuracoes/',
            resultado='NEGADO',
            motivo='Tentativa de aceder ou alterar configurações globais sem privilégios master.'
        )
        return Response(
            {'error': 'Acesso restrito ao Administrador do Sistema. As empresas não têm permissão para aceder à Gestão Global.'},
            status=status.HTTP_403_FORBIDDEN
        )

    import json
    dados = {**DEFAULT_CONFIGURACOES}
    if CONFIG_FILE_PATH.exists():
        try:
            with open(CONFIG_FILE_PATH, 'r', encoding='utf-8') as f:
                salvo = json.load(f)
                if isinstance(salvo, dict):
                    dados.update(salvo)
        except Exception:
            pass

    if request.method in ['POST', 'PUT', 'PATCH']:
        novos_dados = request.data
        for k in DEFAULT_CONFIGURACOES.keys():
            if k in novos_dados:
                dados[k] = novos_dados[k]
        try:
            with open(CONFIG_FILE_PATH, 'w', encoding='utf-8') as f:
                json.dump(dados, f, ensure_ascii=False, indent=2)
        except Exception as e:
            return Response({'error': f'Erro ao persistir configurações: {str(e)}'}, status=500)

    return Response(dados)


@api_view(['POST'])
@permission_classes([AllowAny])
def admin_registrar_tentativa_view(request):
    """
    Endpoint para registar tentativas indevidas de navegação ou acesso ao Admin.
    Permite registar tanto acessos autenticados de utilizadores de empresas como não autenticados.
    """
    rota = request.data.get('rota', '/admin-gestao')
    motivo = request.data.get('motivo', 'Tentativa de navegação direta à Área Administrativa por utilizador de empresa.')
    registrar_tentativa_acesso(request, rota=rota, resultado='NEGADO', motivo=motivo)
    return Response({'status': 'registado', 'message': 'Tentativa de acesso indevido registada para auditoria de segurança.'})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_listar_registros_acesso_view(request):
    """
    Consulta o registo de auditoria de segurança de tentativas de acesso (apenas Super Admin / Staff).
    """
    user = request.user
    is_admin = bool(
        user.is_superuser
        or user.is_staff
        or user.groups.filter(name__in=['Super Admin', 'Administrador']).exists()
    )
    if not is_admin:
        registrar_tentativa_acesso(
            request,
            rota='/api/admin/registros-acesso/',
            resultado='NEGADO',
            motivo='Tentativa de ler registos de auditoria de segurança sem privilégios.'
        )
        return Response({'error': 'Acesso reservado ao Administrador do Sistema.'}, status=status.HTTP_403_FORBIDDEN)

    registros = RegistroAcessoAdmin.objects.all().order_by('-created_at')[:100]
    dados = [
        {
            'id': r.id,
            'username': r.username or 'Anónimo',
            'empresa_nome': r.empresa_nome or 'Não associada',
            'ip_origem': r.ip_origem or '—',
            'user_agent': r.user_agent or '—',
            'rota_tentada': r.rota_tentada,
            'motivo': r.motivo,
            'resultado': r.resultado,
            'created_at': r.created_at.isoformat(),
        }
        for r in registros
    ]
    return Response({
        'total': RegistroAcessoAdmin.objects.count(),
        'total_bloqueados': RegistroAcessoAdmin.objects.filter(resultado='NEGADO').count(),
        'registros': dados
    })


from .pgc_data import (
    CLASSES_PGC,
    CONTAS_PGC,
    PRINCIPIOS_PGC,
    CRITERIOS_VALORIMETRIA,
    DEMONSTRACOES_FINANCEIRAS,
    NOTAS_ANEXO,
    PGC_PME_COMPARATIVO,
    QUADRO_07_RECONCILIACAO_FISCAL,
    DIPLOMA_LEGAL,
)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_pgc_view(request):
    """
    Compêndio Integral do Plano Geral de Contabilidade (PGC) de Angola (Decreto n.º 82/01).
    Retorna o quadro de contas com 157 subcontas normalizadas, princípios contabilísticos,
    critérios de valorimetria, demonstrações financeiras oficiais, 26 notas do anexo,
    comparativo PGC-PME (Dec. Executivo 39/19) e reconciliação fiscal (Quadro 07 Modelo 1 AGT).
    """
    classe_filtro = request.GET.get('classe')
    pesquisa = (request.GET.get('pesquisa') or request.GET.get('q') or '').strip().lower()

    contas = CONTAS_PGC
    if classe_filtro and classe_filtro != 'TODAS':
        try:
            num = int(classe_filtro)
            contas = [c for c in contas if c['classe'] == num]
        except ValueError:
            pass

    if pesquisa:
        contas = [
            c for c in contas
            if pesquisa in c['codigo'].lower()
            or pesquisa in c['nome'].lower()
            or pesquisa in c['fiscal_relevancia'].lower()
            or pesquisa in c.get('debito_por', '').lower()
            or pesquisa in c.get('credito_por', '').lower()
            or pesquisa in c.get('saft_taxonomia', '').lower()
        ]

    return Response({
        'classes': CLASSES_PGC,
        'contas': contas,
        'total': len(contas),
        'total_oficial': len(CONTAS_PGC),
        'principios': PRINCIPIOS_PGC,
        'criterios_valorimetria': CRITERIOS_VALORIMETRIA,
        'demonstracoes_financeiras': DEMONSTRACOES_FINANCEIRAS,
        'notas_anexo': NOTAS_ANEXO,
        'pgc_pme_comparativo': PGC_PME_COMPARATIVO,
        'quadro_07_reconciliacao': QUADRO_07_RECONCILIACAO_FISCAL,
        'diploma_legal': DIPLOMA_LEGAL,
        'base_legal': 'Decreto n.º 82/01, de 16 de Novembro (Diário da República Iª Série, N.º 53)',
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_pgc_export_excel_view(request):
    """
    Gera e faz o download da folha de cálculo Excel (.xlsx) contendo todo o Compêndio
    do Plano Geral de Contabilidade de Angola (Decreto 82/01) estruturado em 7 abas:
    1. Quadro de Contas & Dinâmica (150+ Contas e Taxonomias SAF-T)
    2. Balanço Oficial (Modelo Decreto 82/01)
    3. Demonstração dos Resultados por Natureza (DRN)
    4. Princípios Contabilísticos Fundamentais
    5. Critérios de Valorimetria & Avaliação Patrimonial
    6. Comparativo PGC Geral vs PGC-PME (Dec. Executivo 39/19)
    7. Reconciliação Fiscal: Quadro 07 da Modelo 1 do Imposto Industrial
    """
    import io
    from django.http import HttpResponse
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()

    primary_fill = PatternFill(start_color="5F63F2", end_color="5F63F2", fill_type="solid")
    dark_fill = PatternFill(start_color="272B41", end_color="272B41", fill_type="solid")
    accent_fill = PatternFill(start_color="20C997", end_color="20C997", fill_type="solid")
    warning_fill = PatternFill(start_color="FA8B0C", end_color="FA8B0C", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8F9FE", end_color="F8F9FE", fill_type="solid")

    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    title_font = Font(name="Arial", size=13, bold=True, color="272B41")
    subtitle_font = Font(name="Arial", size=9.5, italic=True, color="6C757D")
    section_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Arial", size=9.5, bold=True)
    regular_font = Font(name="Arial", size=9.5)

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0'),
    )

    # ==============================================================
    # ABA 1: QUADRO DE CONTAS & DINÂMICA CONTABILÍSTICA (150+ CONTAS)
    # ==============================================================
    ws1 = wb.active
    ws1.title = "1. Quadro de Contas (PGC)"

    ws1.merge_cells("A1:J1")
    ws1["A1"] = "REPÚBLICA DE ANGOLA — PLANO GERAL DE CONTABILIDADE (DECRETO N.º 82/01)"
    ws1["A1"].font = title_font
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws1.merge_cells("A2:J2")
    ws1["A2"] = "Quadro de Contas Oficial (157 Subcontas), Dinâmica Débito/Crédito, Mapeamento Fiscal AGT e Taxonomia SAF-T AO"
    ws1["A2"].font = subtitle_font
    ws1["A2"].alignment = Alignment(horizontal="center", vertical="center")

    ws1.row_dimensions[1].height = 26
    ws1.row_dimensions[2].height = 18
    ws1.row_dimensions[4].height = 28

    headers1 = [
        ("Código", 12),
        ("Designação da Conta (PGC)", 38),
        ("Classe Contabilística", 26),
        ("Natureza", 14),
        ("Tipo", 18),
        ("Debita-se por...", 45),
        ("Credita-se por...", 45),
        ("Relevância Fiscal AGT (IVA / II / IRT / INSS)", 48),
        ("Taxonomia SAF-T AO (Dec. Exec. 385/20)", 28),
        ("Base Legal", 30),
    ]

    for col_idx, (text, width) in enumerate(headers1, 1):
        cell = ws1.cell(row=4, column=col_idx, value=text)
        cell.fill = primary_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws1.column_dimensions[col_letter].width = width

    for row_idx, conta in enumerate(CONTAS_PGC, 5):
        ws1.row_dimensions[row_idx].height = 24
        fill = zebra_fill if row_idx % 2 == 0 else None
        vals = [
            conta["codigo"],
            conta["nome"],
            conta["classe_nome"],
            conta["natureza"],
            conta["tipo"],
            conta.get("debito_por", "-"),
            conta.get("credito_por", "-"),
            conta["fiscal_relevancia"],
            conta.get("saft_taxonomia", "-"),
            conta["base_legal"],
        ]
        for col_idx, val in enumerate(vals, 1):
            cell = ws1.cell(row=row_idx, column=col_idx, value=val)
            cell.font = bold_font if col_idx == 1 else regular_font
            cell.border = thin_border
            if fill:
                cell.fill = fill
            if col_idx in [1, 4]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [6, 7, 8]:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            elif col_idx == 9:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # ==============================================================
    # ABA 2: BALANÇO OFICIAL (DECRETO N.º 82/01)
    # ==============================================================
    ws2 = wb.create_sheet(title="2. Balanço Oficial")
    ws2.merge_cells("A1:D1")
    ws2["A1"] = "BALANÇO SINTÉTICO OFICIAL — DECRETO N.º 82/01"
    ws2["A1"].font = title_font
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")

    headers2 = [("Secção / Rubrica Patrimonial", 45), ("Contas PGC Enquadradas", 28), ("Nota no Anexo", 18), ("Tipo", 20)]
    ws2.row_dimensions[3].height = 26
    for col_idx, (text, width) in enumerate(headers2, 1):
        cell = ws2.cell(row=3, column=col_idx, value=text)
        cell.fill = dark_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws2.column_dimensions[col_letter].width = width

    curr_row = 4
    seccoes_balanco = [
        ("ATIVO NÃO CORRENTE (FIXO)", DEMONSTRACOES_FINANCEIRAS["balanco"]["ativo_nao_corrente"], primary_fill),
        ("ATIVO CORRENTE (CIRCULANTE)", DEMONSTRACOES_FINANCEIRAS["balanco"]["ativo_corrente"], primary_fill),
        ("CAPITAL PRÓPRIO (SITUAÇÃO LÍQUIDA)", DEMONSTRACOES_FINANCEIRAS["balanco"]["capital_proprio"], accent_fill),
        ("PASSIVO NÃO CORRENTE (EXIGÍVEL A MÉDIO E LONGO PRAZO)", DEMONSTRACOES_FINANCEIRAS["balanco"]["passivo_nao_corrente"], dark_fill),
        ("PASSIVO CORRENTE (EXIGÍVEL A CURTO PRAZO)", DEMONSTRACOES_FINANCEIRAS["balanco"]["passivo_corrente"], dark_fill),
    ]

    for title, items, fill_color in seccoes_balanco:
        ws2.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=4)
        c = ws2.cell(row=curr_row, column=1, value=title)
        c.fill = fill_color
        c.font = section_font
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws2.row_dimensions[curr_row].height = 22
        curr_row += 1

        for item in items:
            ws2.row_dimensions[curr_row].height = 20
            ws2.cell(row=curr_row, column=1, value=item["rubrica"]).font = bold_font
            ws2.cell(row=curr_row, column=2, value=item["contas"]).alignment = Alignment(horizontal="center")
            ws2.cell(row=curr_row, column=3, value=item["nota"]).alignment = Alignment(horizontal="center")
            ws2.cell(row=curr_row, column=4, value=title.split(" (")[0]).alignment = Alignment(horizontal="center")
            for c_idx in range(1, 5):
                ws2.cell(row=curr_row, column=c_idx).border = thin_border
            curr_row += 1

    # ==============================================================
    # ABA 3: DEMONSTRAÇÃO DOS RESULTADOS POR NATUREZA (DRN)
    # ==============================================================
    ws3 = wb.create_sheet(title="3. Demonstração Resultados DRN")
    ws3.merge_cells("A1:C1")
    ws3["A1"] = "DEMONSTRAÇÃO DOS RESULTADOS POR NATUREZA (DRN) — MODELO DECRETO 82/01"
    ws3["A1"].font = title_font
    ws3["A1"].alignment = Alignment(horizontal="center", vertical="center")

    headers3 = [("Código", 12), ("Designação da Rubrica Oficial", 65), ("Contas do PGC / Fórmula de Apuramento", 35)]
    ws3.row_dimensions[3].height = 26
    for col_idx, (text, width) in enumerate(headers3, 1):
        cell = ws3.cell(row=3, column=col_idx, value=text)
        cell.fill = primary_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws3.column_dimensions[col_letter].width = width

    for row_idx, item in enumerate(DEMONSTRACOES_FINANCEIRAS["drn"]["linhas"], 4):
        ws3.row_dimensions[row_idx].height = 22
        is_subtotal = item["codigo"] in ["MB", "RO", "RC", "RAI", "RLE"]
        fnt = bold_font if is_subtotal else regular_font
        bg = zebra_fill if is_subtotal else None

        c1 = ws3.cell(row=row_idx, column=1, value=item["codigo"])
        c2 = ws3.cell(row=row_idx, column=2, value=item["designacao"])
        c3 = ws3.cell(row=row_idx, column=3, value=item["contas"])

        for c in [c1, c2, c3]:
            c.font = fnt
            c.border = thin_border
            if bg:
                c.fill = bg
        c1.alignment = Alignment(horizontal="center", vertical="center")
        c2.alignment = Alignment(horizontal="left", vertical="center")
        c3.alignment = Alignment(horizontal="center", vertical="center")

    # ==============================================================
    # ABA 4: PRINCÍPIOS CONTABILÍSTICOS FUNDAMENTAIS
    # ==============================================================
    ws4 = wb.create_sheet(title="4. Princípios Contabilísticos")
    ws4.merge_cells("A1:E1")
    ws4["A1"] = "PRINCÍPIOS CONTABILÍSTICOS FUNDAMENTAIS — QUADRO CONCEPTUAL DO PGC"
    ws4["A1"].font = title_font
    ws4["A1"].alignment = Alignment(horizontal="center", vertical="center")

    headers4 = [
        ("Código", 12),
        ("Nome do Princípio", 28),
        ("Conceito Técnico Oficial (Decreto 82/01)", 50),
        ("Exemplo Prático na Empresa", 45),
        ("Impacto em Auditoria & Fiscalidade AGT", 45),
    ]
    ws4.row_dimensions[3].height = 26
    for col_idx, (text, width) in enumerate(headers4, 1):
        cell = ws4.cell(row=3, column=col_idx, value=text)
        cell.fill = accent_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws4.column_dimensions[col_letter].width = width

    for row_idx, princ in enumerate(PRINCIPIOS_PGC, 4):
        ws4.row_dimensions[row_idx].height = 36
        vals = [princ["codigo"], princ["nome"], princ["conceito"], princ["exemplo"], princ["auditoria_relevancia"]]
        for col_idx, val in enumerate(vals, 1):
            cell = ws4.cell(row=row_idx, column=col_idx, value=val)
            cell.font = bold_font if col_idx <= 2 else regular_font
            cell.border = thin_border
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # ==============================================================
    # ABA 5: CRITÉRIOS DE VALORIMETRIA
    # ==============================================================
    ws5 = wb.create_sheet(title="5. Critérios de Valorimetria")
    ws5.merge_cells("A1:D1")
    ws5["A1"] = "CRITÉRIOS DE VALORIMETRIA E AVALIAÇÃO PATRIMONIAL (DECRETO N.º 82/01)"
    ws5["A1"].font = title_font
    ws5["A1"].alignment = Alignment(horizontal="center", vertical="center")

    headers5 = [("Área Patrimonial", 30), ("Enquadramento Normativo", 30), ("Regras de Valorimetria do PGC", 65), ("Impacto Fiscal AGT", 45)]
    ws5.row_dimensions[3].height = 26
    for col_idx, (text, width) in enumerate(headers5, 1):
        cell = ws5.cell(row=3, column=col_idx, value=text)
        cell.fill = dark_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws5.column_dimensions[col_letter].width = width

    curr_row5 = 4
    for crit in CRITERIOS_VALORIMETRIA:
        regras_formatadas = "\n• " + "\n• ".join(crit["regras"])
        ws5.row_dimensions[curr_row5].height = 55
        c1 = ws5.cell(row=curr_row5, column=1, value=crit["area"])
        c2 = ws5.cell(row=curr_row5, column=2, value=crit["norma_pgc"])
        c3 = ws5.cell(row=curr_row5, column=3, value=regras_formatadas)
        c4 = ws5.cell(row=curr_row5, column=4, value=crit["impacto_fiscal"])

        c1.font = bold_font
        c2.font = regular_font
        c3.font = regular_font
        c4.font = bold_font

        for cell in [c1, c2, c3, c4]:
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        curr_row5 += 1

    # ==============================================================
    # ABA 6: COMPARATIVO PGC GERAL VS PGC-PME (DEC. EXECUTIVO 39/19)
    # ==============================================================
    ws6 = wb.create_sheet(title="6. Comparativo PGC vs PME")
    ws6.merge_cells("A1:D1")
    ws6["A1"] = "QUADRO COMPARATIVO: PGC GERAL (DEC. 82/01) vs PGC-PME (DEC. EXECUTIVO 39/19)"
    ws6["A1"].font = title_font
    ws6["A1"].alignment = Alignment(horizontal="center", vertical="center")

    headers6 = [("Critério de Comparação", 30), ("PGC Geral (Decreto 82/01)", 45), ("PGC-PME (Decreto Executivo 39/19)", 45), ("Relevância Prática no Escritório", 45)]
    ws6.row_dimensions[3].height = 26
    for col_idx, (text, width) in enumerate(headers6, 1):
        cell = ws6.cell(row=3, column=col_idx, value=text)
        cell.fill = accent_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws6.column_dimensions[col_letter].width = width

    for row_idx, item in enumerate(PGC_PME_COMPARATIVO, 4):
        ws6.row_dimensions[row_idx].height = 36
        c1 = ws6.cell(row=row_idx, column=1, value=item["criterio"])
        c2 = ws6.cell(row=row_idx, column=2, value=item["pgc_geral"])
        c3 = ws6.cell(row=row_idx, column=3, value=item["pgc_pme"])
        c4 = ws6.cell(row=row_idx, column=4, value=item["relevancia"])

        c1.font = bold_font
        c2.font = regular_font
        c3.font = regular_font
        c4.font = regular_font

        for c in [c1, c2, c3, c4]:
            c.border = thin_border
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # ==============================================================
    # ABA 7: RECONCILIAÇÃO FISCAL: QUADRO 07 DA MODELO 1 DA AGT
    # ==============================================================
    ws7 = wb.create_sheet(title="7. Reconciliação Fiscal Q07")
    ws7.merge_cells("A1:E1")
    ws7["A1"] = "RECONCILIAÇÃO FISCAL: QUADRO 07 DA DECLARAÇÃO MODELO 1 DO IMPOSTO INDUSTRIAL"
    ws7["A1"].font = title_font
    ws7["A1"].alignment = Alignment(horizontal="center", vertical="center")

    headers7 = [("Campo Modelo 1", 15), ("Descrição da Correção Fiscal", 50), ("Conta PGC", 24), ("Operação", 35), ("Fundamento Legal (Lei 19/14)", 40)]
    ws7.row_dimensions[3].height = 26
    for col_idx, (text, width) in enumerate(headers7, 1):
        cell = ws7.cell(row=3, column=col_idx, value=text)
        cell.fill = warning_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws7.column_dimensions[col_letter].width = width

    for row_idx, q in enumerate(QUADRO_07_RECONCILIACAO_FISCAL, 4):
        ws7.row_dimensions[row_idx].height = 30
        c1 = ws7.cell(row=row_idx, column=1, value=q["campo"])
        c2 = ws7.cell(row=row_idx, column=2, value=q["descricao"])
        c3 = ws7.cell(row=row_idx, column=3, value=q["conta_pgc"])
        c4 = ws7.cell(row=row_idx, column=4, value=q["operacao"])
        c5 = ws7.cell(row=row_idx, column=5, value=q["fundamento_legal"])

        c1.font = bold_font
        c2.font = bold_font if "MATÉRIA" in q["descricao"] else regular_font
        c3.font = regular_font
        c4.font = bold_font
        c5.font = regular_font

        for c in [c1, c2, c3, c4, c5]:
            c.border = thin_border
            c.alignment = Alignment(horizontal="left", vertical="center")
        c1.alignment = Alignment(horizontal="center", vertical="center")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="PGC_Angola_Decreto_82_01_Completo.xlsx"'
    return response


# ==============================================================================
# MÓDULO CONTABILÍSTICO OPERACIONAL (DIÁRIO, RAZÃO E BALANCETE GERADO VIA PGC)
# ==============================================================================

@api_view(['GET', 'POST'])
@permission_classes([IsAuthenticated])
def listar_criar_lancamentos_view(request):
    """
    Livro Diário de Lançamentos Contabilísticos (Partidas Dobradas).
    GET: Lista lançamentos filtrados por empresa, período e tipo de diário.
    POST: Regista novo lançamento garantindo Débito == Crédito e conformidade PGC.
    """
    user = request.user
    acessiveis = empresas_acessiveis(user)

    if request.method == 'GET':
        qs = LancamentoContabilistico.objects.filter(empresa__in=acessiveis).select_related('empresa', 'criado_por').prefetch_related('linhas')

        empresa_id = _inteiro_opcional(request.query_params.get('empresa') or request.query_params.get('empresa_id'))
        if empresa_id:
            qs = qs.filter(empresa_id=empresa_id)

        ano = _inteiro_opcional(request.query_params.get('ano'))
        if ano:
            qs = qs.filter(ano=ano)

        mes = _inteiro_opcional(request.query_params.get('mes'))
        if mes:
            qs = qs.filter(mes=mes)

        tipo_diario = request.query_params.get('tipo_diario')
        if tipo_diario and tipo_diario != 'TODOS':
            qs = qs.filter(tipo_diario=tipo_diario)

        estado = request.query_params.get('estado')
        if estado and estado != 'TODOS':
            qs = qs.filter(estado=estado)

        pesquisa = (request.query_params.get('pesquisa') or request.query_params.get('q') or '').strip()
        if pesquisa:
            qs = qs.filter(
                Q(numero_diario__icontains=pesquisa)
                | Q(descricao__icontains=pesquisa)
                | Q(documento_referencia__icontains=pesquisa)
                | Q(linhas__terceiro_nome__icontains=pesquisa)
                | Q(linhas__terceiro_nif__icontains=pesquisa)
                | Q(linhas__conta_codigo__icontains=pesquisa)
                | Q(linhas__conta_nome__icontains=pesquisa)
            ).distinct()

        serializer = LancamentoContabilisticoSerializer(qs, many=True)
        return Response(serializer.data)

    elif request.method == 'POST':
        data = request.data.copy() if hasattr(request.data, 'copy') else dict(request.data)
        empresa_id = _inteiro_opcional(data.get('empresa') or data.get('empresa_id'))
        if not empresa_id or not acessiveis.filter(pk=empresa_id).exists():
            return Response({'error': 'Empresa inválida ou sem permissão de acesso.'}, status=status.HTTP_400_BAD_REQUEST)

        data['empresa'] = empresa_id
        ano = int(data.get('ano') or 2026)
        mes = int(data.get('mes') or 7)
        data['ano'] = ano
        data['mes'] = mes

        if not data.get('data_lancamento'):
            data['data_lancamento'] = f"{ano}-{mes:02d}-01"

        if not data.get('numero_diario'):
            ultimo = LancamentoContabilistico.objects.filter(empresa_id=empresa_id, ano=ano).count()
            data['numero_diario'] = f"DIA-{ano}-{(ultimo + 1):04d}"

        # Calcular totais a partir das linhas
        linhas = data.get('linhas', [])
        total_d = sum(Decimal(str(l.get('debito', 0) or 0)) for l in linhas)
        total_c = sum(Decimal(str(l.get('credito', 0) or 0)) for l in linhas)
        data['total_debito'] = str(total_d)
        data['total_credito'] = str(total_c)

        serializer = LancamentoContabilisticoSerializer(data=data)
        if serializer.is_valid():
            lancamento = serializer.save(criado_por=user, empresa_id=empresa_id)
            return Response(LancamentoContabilisticoSerializer(lancamento).data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


@api_view(['GET', 'PUT', 'PATCH', 'DELETE'])
@permission_classes([IsAuthenticated])
def detalhe_lancamento_view(request, lancamento_id):
    """Consulta, edição ou anulação de um lançamento contabilístico específico."""
    acessiveis = empresas_acessiveis(request.user)
    lancamento = LancamentoContabilistico.objects.filter(pk=lancamento_id, empresa__in=acessiveis).first()
    if not lancamento:
        return Response({'error': 'Lançamento não encontrado ou sem permissão.'}, status=status.HTTP_404_NOT_FOUND)

    if request.method == 'GET':
        return Response(LancamentoContabilisticoSerializer(lancamento).data)

    elif request.method in ['PUT', 'PATCH']:
        serializer = LancamentoContabilisticoSerializer(lancamento, data=request.data, partial=(request.method == 'PATCH'))
        if serializer.is_valid():
            # Se vierem linhas novas na edição, atualizar
            if 'linhas' in request.data:
                lancamento.linhas.all().delete()
                for idx, l in enumerate(request.data.get('linhas', []), start=1):
                    LinhaLancamento.objects.create(
                        lancamento=lancamento,
                        ordem=l.get('ordem', idx),
                        conta_codigo=l.get('conta_codigo', '').strip(),
                        conta_nome=l.get('conta_nome', '').strip(),
                        debito=Decimal(str(l.get('debito', 0) or 0)),
                        credito=Decimal(str(l.get('credito', 0) or 0)),
                        descricao_linha=l.get('descricao_linha', '').strip(),
                        terceiro_nif=l.get('terceiro_nif', '').strip(),
                        terceiro_nome=l.get('terceiro_nome', '').strip(),
                    )
            lancamento = serializer.save()
            return Response(LancamentoContabilisticoSerializer(lancamento).data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    elif request.method == 'DELETE':
        lancamento.delete()
        return Response({'success': True, 'message': 'Lançamento eliminado com sucesso.'}, status=status.HTTP_200_OK)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def razao_contabilistico_view(request):
    """
    Livro Razão / Extrato de Contas do PGC.
    Permite inspecionar todos os movimentos a débito e a crédito de uma conta
    (ex: 43.1 Bancos, 31.1 Clientes, 34.5 IVA) e calcular o saldo acumulado.
    """
    acessiveis = empresas_acessiveis(request.user)
    empresa_id = _inteiro_opcional(request.query_params.get('empresa') or request.query_params.get('empresa_id'))
    if not empresa_id or not acessiveis.filter(pk=empresa_id).exists():
        return Response({'error': 'Empresa inválida ou não especificada.'}, status=status.HTTP_400_BAD_REQUEST)

    ano = _inteiro_opcional(request.query_params.get('ano')) or 2026
    mes = _inteiro_opcional(request.query_params.get('mes'))
    conta_codigo = (request.query_params.get('conta') or request.query_params.get('conta_codigo') or '').strip()

    linhas_qs = LinhaLancamento.objects.filter(
        lancamento__empresa_id=empresa_id,
        lancamento__ano=ano,
        lancamento__estado='CONFIRMADO'
    ).select_related('lancamento').order_by('lancamento__data_lancamento', 'lancamento__id', 'ordem')

    if conta_codigo:
        linhas_qs = linhas_qs.filter(conta_codigo__startswith=conta_codigo)

    if mes:
        linhas_qs = linhas_qs.filter(lancamento__mes=mes)

    # Agrupar por conta
    contas_map = {}
    for linha in linhas_qs:
        cod = linha.conta_codigo
        if cod not in contas_map:
            contas_map[cod] = {
                'conta_codigo': cod,
                'conta_nome': linha.conta_nome,
                'total_debito': Decimal('0.00'),
                'total_credito': Decimal('0.00'),
                'saldo_final': Decimal('0.00'),
                'natureza_saldo': 'Devedor',
                'movimentos': [],
            }

        item = contas_map[cod]
        item['total_debito'] += linha.debito
        item['total_credito'] += linha.credito

        saldo_parcial = item['total_debito'] - item['total_credito']
        item['saldo_final'] = abs(saldo_parcial)
        item['natureza_saldo'] = 'Devedor' if saldo_parcial >= 0 else 'Credor'

        item['movimentos'].append({
            'lancamento_id': linha.lancamento.id,
            'numero_diario': linha.lancamento.numero_diario,
            'data': linha.lancamento.data_lancamento.strftime('%Y-%m-%d'),
            'tipo_diario': linha.lancamento.tipo_diario,
            'descricao': linha.descricao_linha or linha.lancamento.descricao,
            'documento_referencia': linha.lancamento.documento_referencia,
            'terceiro_nome': linha.terceiro_nome,
            'terceiro_nif': linha.terceiro_nif,
            'debito': str(linha.debito),
            'credito': str(linha.credito),
            'saldo_apos': str(abs(saldo_parcial)),
            'natureza_apos': item['natureza_saldo'],
        })

    lista_contas = sorted(contas_map.values(), key=lambda x: x['conta_codigo'])
    return Response({
        'empresa_id': empresa_id,
        'ano': ano,
        'mes': mes,
        'total_contas_movimentadas': len(lista_contas),
        'contas': lista_contas,
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def balancete_gerado_view(request):
    """
    Geração Oficial do Balancete de Verificação a partir dos Lançamentos Reais da Empresa.
    Obedece às 8 Classes do PGC (Decreto n.º 82/01):
    Calcula Débitos, Créditos, Saldo Devedor e Saldo Credor por Conta e por Classe.
    """
    acessiveis = empresas_acessiveis(request.user)
    empresa_id = _inteiro_opcional(request.query_params.get('empresa') or request.query_params.get('empresa_id'))
    if not empresa_id or not acessiveis.filter(pk=empresa_id).exists():
        return Response({'error': 'Empresa não encontrada ou sem permissão.'}, status=status.HTTP_400_BAD_REQUEST)

    ano = _inteiro_opcional(request.query_params.get('ano')) or 2026
    mes = _inteiro_opcional(request.query_params.get('mes')) or 7

    empresa = acessiveis.get(pk=empresa_id)

    # Buscar todos os lançamentos confirmados até ao mês selecionado
    linhas = LinhaLancamento.objects.filter(
        lancamento__empresa_id=empresa_id,
        lancamento__ano=ano,
        lancamento__mes__lte=mes,
        lancamento__estado='CONFIRMADO'
    ).select_related('lancamento')

    contas_balancete = {}
    totais = {
        'total_debito_mes': Decimal('0.00'),
        'total_credito_mes': Decimal('0.00'),
        'total_debito_acumulado': Decimal('0.00'),
        'total_credito_acumulado': Decimal('0.00'),
        'total_saldo_devedor': Decimal('0.00'),
        'total_saldo_credor': Decimal('0.00'),
    }

    for l in linhas:
        cod = l.conta_codigo
        if cod not in contas_balancete:
            classe_num = int(cod[0]) if cod and cod[0].isdigit() else 0
            contas_balancete[cod] = {
                'codigo': cod,
                'nome': l.conta_nome,
                'classe': classe_num,
                'debito_mes': Decimal('0.00'),
                'credito_mes': Decimal('0.00'),
                'debito_acumulado': Decimal('0.00'),
                'credito_acumulado': Decimal('0.00'),
                'saldo_devedor': Decimal('0.00'),
                'saldo_credor': Decimal('0.00'),
            }

        c = contas_balancete[cod]
        c['debito_acumulado'] += l.debito
        c['credito_acumulado'] += l.credito

        if l.lancamento.mes == mes:
            c['debito_mes'] += l.debito
            c['credito_mes'] += l.credito

    # Calcular saldos e totais gerais
    linhas_lista = []
    for cod in sorted(contas_balancete.keys()):
        c = contas_balancete[cod]
        saldo = c['debito_acumulado'] - c['credito_acumulado']
        if saldo > 0:
            c['saldo_devedor'] = saldo
            c['saldo_credor'] = Decimal('0.00')
        elif saldo < 0:
            c['saldo_devedor'] = Decimal('0.00')
            c['saldo_credor'] = abs(saldo)

        totais['total_debito_mes'] += c['debito_mes']
        totais['total_credito_mes'] += c['credito_mes']
        totais['total_debito_acumulado'] += c['debito_acumulado']
        totais['total_credito_acumulado'] += c['credito_acumulado']
        totais['total_saldo_devedor'] += c['saldo_devedor']
        totais['total_saldo_credor'] += c['saldo_credor']

        linhas_lista.append({
            'codigo': c['codigo'],
            'nome': c['nome'],
            'classe': c['classe'],
            'debito_mes': str(c['debito_mes']),
            'credito_mes': str(c['credito_mes']),
            'debito_acumulado': str(c['debito_acumulado']),
            'credito_acumulado': str(c['credito_acumulado']),
            'saldo_devedor': str(c['saldo_devedor']),
            'saldo_credor': str(c['saldo_credor']),
        })

    balanceado = (totais['total_saldo_devedor'] == totais['total_saldo_credor'])

    return Response({
        'empresa_id': empresa.id,
        'empresa_nome': empresa.nome,
        'empresa_nif': empresa.nif,
        'ano': ano,
        'mes': mes,
        'balanceado': balanceado,
        'totais': {k: str(v) for k, v in totais.items()},
        'total_contas': len(linhas_lista),
        'linhas': linhas_lista,
    })


# ==============================================================================
# LEGISLAÇÃO FISCAL DE ANGOLA - COMPÊNDIO & EXPORTAÇÃO EXCEL
# ==============================================================================

@api_view(['GET'])
@permission_classes([AllowAny])
def legislacao_fiscal_view(request):
    """
    Compêndio Oficial da Legislação Fiscal Angolana.
    Suporta filtros por código de diploma (CIVA, CII, CIRT, CIS, CGT, RJFDE, INSS)
    e pesquisa textual em tempo real sobre artigos, títulos e âmbito.
    """
    from .legislacao_fiscal_data import get_diplomas_fiscais, DIPLOMAS_FISCAIS
    filtro_codigo = request.query_params.get('diploma') or request.query_params.get('codigo')
    termo = request.query_params.get('pesquisa') or request.query_params.get('q') or ''

    diplomas = get_diplomas_fiscais(filtro_codigo, termo)
    return Response({
        'total_diplomas': len(diplomas),
        'diplomas': diplomas,
        'codigos_disponiveis': [
            {'codigo': d['codigo'], 'nome': d['nome'], 'taxa_geral': d['taxa_geral']}
            for d in DIPLOMAS_FISCAIS
        ]
    })


@api_view(['GET'])
@permission_classes([AllowAny])
def legislacao_fiscal_export_excel_view(request):
    """
    Gera e faz o download da folha de cálculo Excel (.xlsx) contendo todo o Compêndio
    da Legislação Fiscal de Angola estruturado por Diplomas e Artigos Fiscais.
    """
    import io
    from django.http import HttpResponse
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from .legislacao_fiscal_data import DIPLOMAS_FISCAIS

    wb = openpyxl.Workbook()

    primary_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    dark_fill = PatternFill(start_color="272B41", end_color="272B41", fill_type="solid")
    accent_fill = PatternFill(start_color="2B50ED", end_color="2B50ED", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8F9FE", end_color="F8F9FE", fill_type="solid")

    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    title_font = Font(name="Arial", size=13, bold=True, color="272B41")
    subtitle_font = Font(name="Arial", size=9.5, italic=True, color="6C757D")
    section_font = Font(name="Arial", size=10.5, bold=True, color="FFFFFF")
    bold_font = Font(name="Arial", size=9.5, bold=True)
    regular_font = Font(name="Arial", size=9.5)

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # ABA 1: RESUMO DOS DIPLOMAS FISCAIS
    ws1 = wb.active
    ws1.title = "Códigos Tributários Angola"
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells('A1:E1')
    ws1['A1'] = "COMPÊNDIO OFICIAL DA LEGISLAÇÃO FISCAL DE ANGOLA"
    ws1['A1'].font = title_font
    ws1['A1'].alignment = Alignment(vertical="center")
    ws1.row_dimensions[1].height = 28

    ws1.merge_cells('A2:E2')
    ws1['A2'] = "Enquadramento Legal, Taxas, Prazos Declarativos e Normas Tributárias em Vigor"
    ws1['A2'].font = subtitle_font
    ws1.row_dimensions[2].height = 18

    headers1 = ["Código", "Diploma Fiscal", "Base Legal", "Taxa Geral / Regime", "Âmbito de Aplicação"]
    for col_idx, h in enumerate(headers1, 1):
        cell = ws1.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = primary_fill
        cell.alignment = Alignment(horizontal="center" if col_idx == 1 else "left", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws1.row_dimensions[4].height = 26

    current_row = 5
    for d in DIPLOMAS_FISCAIS:
        ws1.cell(row=current_row, column=1, value=d['codigo']).alignment = Alignment(horizontal="center", vertical="top")
        ws1.cell(row=current_row, column=2, value=d['nome']).alignment = Alignment(vertical="top", wrap_text=True)
        ws1.cell(row=current_row, column=3, value=d['diploma']).alignment = Alignment(vertical="top", wrap_text=True)
        ws1.cell(row=current_row, column=4, value=d['taxa_geral']).alignment = Alignment(vertical="top", wrap_text=True)
        ws1.cell(row=current_row, column=5, value=d['ambito']).alignment = Alignment(vertical="top", wrap_text=True)

        for col_idx in range(1, 6):
            cell = ws1.cell(row=current_row, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            if current_row % 2 == 0:
                cell.fill = zebra_fill
        ws1.row_dimensions[current_row].height = 42
        current_row += 1

    ws1.column_dimensions['A'].width = 12
    ws1.column_dimensions['B'].width = 38
    ws1.column_dimensions['C'].width = 36
    ws1.column_dimensions['D'].width = 24
    ws1.column_dimensions['E'].width = 50

    # ABA 2: ARTIGOS E NORMAS DETALHADAS POR DIPLOMA
    ws2 = wb.create_sheet(title="Artigos e Normas Chave")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells('A1:D1')
    ws2['A1'] = "ARTIGOS-CHAVE E NORMAS TRIBUTÁRIAS POR DIPLOMA FISCAL"
    ws2['A1'].font = title_font
    ws2['A1'].alignment = Alignment(vertical="center")
    ws2.row_dimensions[1].height = 28

    headers2 = ["Diploma", "Artigo / Norma", "Título / Matéria", "Disposição Legal e Procedimento"]
    for col_idx, h in enumerate(headers2, 1):
        cell = ws2.cell(row=3, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = accent_fill
        cell.alignment = Alignment(horizontal="center" if col_idx <= 2 else "left", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws2.row_dimensions[3].height = 26

    current_row = 4
    for d in DIPLOMAS_FISCAIS:
        for a in d['artigos_chave']:
            ws2.cell(row=current_row, column=1, value=d['codigo']).alignment = Alignment(horizontal="center", vertical="top")
            ws2.cell(row=current_row, column=2, value=a['artigo']).alignment = Alignment(horizontal="center", vertical="top")
            ws2.cell(row=current_row, column=3, value=a['titulo']).alignment = Alignment(vertical="top", wrap_text=True)
            ws2.cell(row=current_row, column=4, value=a['conteudo']).alignment = Alignment(vertical="top", wrap_text=True)

            for col_idx in range(1, 5):
                cell = ws2.cell(row=current_row, column=col_idx)
                cell.font = regular_font
                cell.border = thin_border
                if current_row % 2 == 0:
                    cell.fill = zebra_fill
            ws2.row_dimensions[current_row].height = 50
            current_row += 1

    ws2.column_dimensions['A'].width = 12
    ws2.column_dimensions['B'].width = 20
    ws2.column_dimensions['C'].width = 34
    ws2.column_dimensions['D'].width = 75

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="Legislacao_Fiscal_Angola_Compendio.xlsx"'
    return response


@api_view(['GET', 'POST'])
@permission_classes([AllowAny])
def confronto_export_excel_view(request, analise_id=None):
    """
    Gera e faz o download da folha de cálculo Excel (.xlsx) oficial contendo o relatório
    de confronto e reconciliação entre o Balancete (PGC) e o Modelo 7 da AGT.
    """
    from datetime import datetime
    from django.http import HttpResponse
    from .confronto_excel import gerar_excel_confronto

    texto = ""
    empresa_nome = "Empresa"
    nif = "---"
    periodo = "---"
    doc_contab = "Balancete de Verificação (PGC)"
    doc_agt = "Declaração Modelo 7 da AGT"

    if analise_id:
        analise = AnaliseIA.objects.select_related(
            'documento_contabilidade__empresa', 'documento_agt__empresa'
        ).filter(pk=analise_id).first()
        if not analise:
            return Response({'error': 'Análise não encontrada'}, status=status.HTTP_404_NOT_FOUND)
        texto = analise.resultado or ""
        doc_c = analise.documento_contabilidade
        doc_a = analise.documento_agt
        if doc_c and doc_c.empresa:
            empresa_nome = doc_c.empresa.nome
            nif = doc_c.empresa.nif
            periodo = f"{doc_c.mes:02d}/{doc_c.ano}" if doc_c.mes and doc_c.ano else "---"
        elif doc_a and doc_a.empresa:
            empresa_nome = doc_a.empresa.nome
            nif = doc_a.empresa.nif
            periodo = f"{doc_a.mes:02d}/{doc_a.ano}" if doc_a.mes and doc_a.ano else "---"
        if doc_c:
            doc_contab = doc_c.nome_arquivo
        if doc_a:
            doc_agt = doc_a.nome_arquivo
    else:
        texto = request.data.get('texto') or request.query_params.get('texto') or ""
        empresa_nome = request.data.get('empresa_nome') or request.query_params.get('empresa_nome') or empresa_nome
        nif = request.data.get('nif') or request.query_params.get('nif') or nif
        periodo = request.data.get('periodo') or request.query_params.get('periodo') or periodo
        doc_contab = request.data.get('doc_contab') or request.query_params.get('doc_contab') or doc_contab
        doc_agt = request.data.get('doc_agt') or request.query_params.get('doc_agt') or doc_agt

    if not texto:
        return Response({'error': 'Texto da análise é obrigatório para gerar o relatório.'}, status=status.HTTP_400_BAD_REQUEST)

    excel_bytes = gerar_excel_confronto(
        texto_analise=texto,
        empresa_nome=empresa_nome,
        nif=nif,
        periodo=periodo,
        doc_contab=doc_contab,
        doc_agt=doc_agt
    )

    filename = f"Relatorio_Confronto_Fiscal_{nif}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(
        excel_bytes,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@api_view(['GET'])
@permission_classes([AllowAny])
def consolidacao_anual_view(request, empresa_id, ano):
    """Retorna o confronto consolidado de 12 meses de Modelo 7 com o Balancete Anual de Fecho."""
    try:
        resultado = consolidacao_anual.gerar_consolidacao_anual(int(empresa_id), int(ano))
        return Response(resultado, status=status.HTTP_200_OK)
    except Empresa.DoesNotExist:
        return Response({'error': 'Empresa não encontrada.'}, status=status.HTTP_404_NOT_FOUND)
    except Exception as e:
        logger.exception(f"Erro ao gerar consolidação anual: {e}")
        return Response({'error': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


@api_view(['GET'])
@permission_classes([AllowAny])
def consolidacao_anual_exportar_excel_view(request, empresa_id, ano):
    """Exporta o livro Excel (.xlsx) da consolidação anual dos 12 meses."""
    try:
        resultado = consolidacao_anual.gerar_consolidacao_anual(int(empresa_id), int(ano))
        excel_bytes = consolidacao_anual.gerar_excel_consolidacao_anual(resultado)
        nif = resultado['empresa']['nif']
        filename = f"Consolidacao_Anual_Fiscal_{nif}_{ano}.xlsx"
        response = HttpResponse(
            excel_bytes,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
    except Empresa.DoesNotExist:
        return Response({'error': 'Empresa não encontrada.'}, status=status.HTTP_404_NOT_FOUND)
    except Exception as e:
        logger.exception(f"Erro ao exportar Excel da consolidação anual: {e}")
        return Response({'error': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


# ============================================================================
# MÓDULO INTELIGENTE: LEITOR / SCANNER DE FACTURAS COM IA & AUDITORIA AGT
# ============================================================================

@api_view(['POST'])
@permission_classes([AllowAny])
def scan_factura_view(request):
    """
    Recebe um ficheiro de fatura (PDF, imagem, texto), processa com IA especializada
    nas regras fiscais de Angola (AGT/CIVA/PGC), audita e grava na base de dados.
    """
    arquivo = request.FILES.get('arquivo')
    if not arquivo:
        return Response({'error': 'Ficheiro da fatura não fornecido.'}, status=status.HTTP_400_BAD_REQUEST)

    empresa_id = request.data.get('empresa_id') or request.query_params.get('empresa_id')
    import tempfile
    import os

    extensao = ('.' + arquivo.name.rsplit('.', 1)[-1].lower() if '.' in arquivo.name else '')
    with tempfile.NamedTemporaryFile(suffix=extensao, delete=False) as tmp:
        for chunk in arquivo.chunks():
            tmp.write(chunk)
        tmp_path = tmp.name

    try:
        empresa_obj = None
        if empresa_id:
            empresa_obj = Empresa.objects.filter(pk=empresa_id).first()

        # Arquivar também em DocumentoUpload
        tipo_upload = 'FACTURA_COMPRA'
        if 'venda' in arquivo.name.lower() or 'cliente' in arquivo.name.lower():
            tipo_upload = 'FACTURA_VENDA'

        doc_upload = DocumentoUpload.objects.create(
            empresa=empresa_obj,
            arquivo=arquivo,
            tipo=tipo_upload,
            tipo_sugerido=tipo_upload,
            nome_arquivo=arquivo.name,
            estado='PROCESSADO'
        )

        factura = factura_ai_service.processar_factura_ficheiro(
            arquivo_path=tmp_path,
            nome_arquivo=arquivo.name,
            empresa_id=empresa_obj.id if empresa_obj else None,
            doc_upload_id=doc_upload.id
        )

        # Se o DocumentoUpload não tinha empresa ou período, atualiza agora
        if not doc_upload.empresa and factura.empresa:
            doc_upload.empresa = factura.empresa
        doc_upload.ano = factura.ano
        doc_upload.mes = factura.mes
        doc_upload.save(update_fields=['empresa', 'ano', 'mes'])

        serializer = FacturaSerializer(factura)
        return Response({
            'success': True,
            'mensagem': f"Fatura {factura.numero_factura} processada e auditada com sucesso com IA.",
            'factura': serializer.data
        }, status=status.HTTP_201_CREATED)

    except Exception as e:
        logger.exception(f"Erro ao processar fatura com IA: {e}")
        return Response({
            'error': f"Falha ao ler fatura: {str(e)}"
        }, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
    finally:
        if os.path.exists(tmp_path):
            try:
                os.unlink(tmp_path)
            except OSError:
                pass


@api_view(['GET'])
@permission_classes([AllowAny])
def listar_facturas_view(request):
    """
    Lista facturas filtradas por empresa, período e conformidade,
    acompanhadas de métricas fiscais acumuladas para o mapa de IVA.
    """
    empresa_id = request.query_params.get('empresa_id')
    ano = request.query_params.get('ano')
    mes = request.query_params.get('mes')
    tipo_operacao = request.query_params.get('tipo_operacao')
    conformidade = request.query_params.get('conformidade')
    busca = request.query_params.get('search', '').strip()

    qs = Factura.objects.select_related('empresa', 'documento', 'lancamento').all()

    if empresa_id:
        qs = qs.filter(empresa_id=empresa_id)
    if ano:
        try:
            qs = qs.filter(ano=int(ano))
        except ValueError:
            pass
    if mes and mes != 'todos':
        try:
            qs = qs.filter(mes=int(mes))
        except ValueError:
            pass
    if tipo_operacao:
        qs = qs.filter(tipo_operacao=tipo_operacao)
    if conformidade and conformidade != 'todos':
        qs = qs.filter(conformidade_agt=conformidade)
    if busca:
        qs = qs.filter(
            Q(numero_factura__icontains=busca) |
            Q(fornecedor_nome__icontains=busca) |
            Q(fornecedor_nif__icontains=busca) |
            Q(cliente_nome__icontains=busca) |
            Q(cliente_nif__icontains=busca)
        )

    # Estatísticas e Métricas de IVA do lote
    total_base = Decimal('0.00')
    total_iva = Decimal('0.00')
    total_iva_dedutivel = Decimal('0.00')
    total_retencao_iva = Decimal('0.00')
    total_retencao_ii = Decimal('0.00')
    total_bruto = Decimal('0.00')
    total_liquido = Decimal('0.00')
    conformes = 0
    avisos = 0
    irregulares = 0
    lancadas = 0

    facturas_list = list(qs)
    for f in facturas_list:
        total_base += f.base_incidencia or Decimal('0.00')
        total_iva += f.valor_iva or Decimal('0.00')
        if f.iva_dedutivel:
            total_iva_dedutivel += f.valor_iva or Decimal('0.00')
        total_retencao_iva += f.retencao_iva or Decimal('0.00')
        total_retencao_ii += f.retencao_ii or Decimal('0.00')
        total_bruto += f.valor_total or Decimal('0.00')
        total_liquido += f.valor_liquido or Decimal('0.00')

        if f.conformidade_agt == 'CONFORME':
            conformes += 1
        elif f.conformidade_agt == 'AVISO':
            avisos += 1
        else:
            irregulares += 1

        if f.lancado_na_contabilidade:
            lancadas += 1

    serializer = FacturaSerializer(facturas_list, many=True)
    return Response({
        'facturas': serializer.data,
        'estatisticas': {
            'total_contagem': len(facturas_list),
            'total_base': float(total_base),
            'total_iva_suportado': float(total_iva),
            'total_iva_dedutivel': float(total_iva_dedutivel),
            'total_retencao_iva': float(total_retencao_iva),
            'total_retencao_ii': float(total_retencao_ii),
            'total_bruto': float(total_bruto),
            'total_liquido': float(total_liquido),
            'conformes': conformes,
            'avisos': avisos,
            'irregulares': irregulares,
            'lancadas': lancadas,
        }
    })


@api_view(['POST'])
@permission_classes([AllowAny])
def lancar_factura_contabilidade_view(request, factura_id):
    """
    Converte uma fatura auditada num lançamento oficial em partidas dobradas
    no Diário de Compras do PGC de Angola com 1 clique.
    """
    try:
        factura = Factura.objects.get(pk=factura_id)
        user = request.user if getattr(request, 'user', None) and request.user.is_authenticated else None
        lancamento, criado, mensagem = factura_ai_service.criar_lancamento_contabilistico_da_factura(factura, user=user)

        serializer = FacturaSerializer(factura)
        return Response({
            'success': True,
            'mensagem': mensagem,
            'lancamento_id': lancamento.id if lancamento else None,
            'numero_diario': lancamento.numero_diario if lancamento else None,
            'factura': serializer.data
        }, status=status.HTTP_200_OK if not criado else status.HTTP_201_CREATED)
    except Factura.DoesNotExist:
        return Response({'error': 'Fatura não encontrada.'}, status=status.HTTP_404_NOT_FOUND)
    except Exception as e:
        logger.exception(f"Erro ao lançar fatura no diário: {e}")
        return Response({'error': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


@api_view(['DELETE'])
@permission_classes([AllowAny])
def deletar_factura_view(request, factura_id):
    """Elimina o registo da fatura da base de dados."""
    try:
        factura = Factura.objects.get(pk=factura_id)
        numero = factura.numero_factura
        factura.delete()
        return Response({'success': True, 'mensagem': f"Fatura {numero} eliminada com sucesso."})
    except Factura.DoesNotExist:
        return Response({'error': 'Fatura não encontrada.'}, status=status.HTTP_404_NOT_FOUND)
    except Exception as e:
        return Response({'error': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


@api_view(['GET'])
@permission_classes([AllowAny])
def exportar_facturas_excel_view(request):
    """
    Gera o Mapa de Fornecedores / Aquisições da AGT em formato Excel (.xlsx).
    Elemento oficial exigido como anexo à Declaração Periódica de IVA (Modelo 7).
    """
    empresa_id = request.query_params.get('empresa_id')
    ano = request.query_params.get('ano')
    mes = request.query_params.get('mes')

    qs = Factura.objects.select_related('empresa').all()
    if empresa_id:
        qs = qs.filter(empresa_id=empresa_id)
    if ano:
        try:
            qs = qs.filter(ano=int(ano))
        except ValueError:
            pass
    if mes and mes != 'todos':
        try:
            qs = qs.filter(mes=int(mes))
        except ValueError:
            pass

    empresa = Empresa.objects.filter(pk=empresa_id).first() if empresa_id else None
    empresa_nome = empresa.nome if empresa else "Todas as Empresas"
    empresa_nif = empresa.nif if empresa else "---"

    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Mapa de Fornecedores AGT"
    ws.views.sheetView[0].showGridLines = True

    # Estilos
    cor_azul = "1e3a8a"
    cor_topo = "0f172a"
    cor_linha_clara = "f8fafc"
    cor_borda = "cbd5e1"

    fonte_titulo = Font(name="Calibri", size=15, bold=True, color="ffffff")
    fonte_sub = Font(name="Calibri", size=10, bold=False, color="e2e8f0")
    fonte_cabecalho = Font(name="Calibri", size=10, bold=True, color="ffffff")
    fonte_dados = Font(name="Calibri", size=9.5)
    fonte_totais = Font(name="Calibri", size=10, bold=True, color="0f172a")

    fill_titulo = PatternFill(start_color=cor_topo, end_color=cor_topo, fill_type="solid")
    fill_cabecalho = PatternFill(start_color=cor_azul, end_color=cor_azul, fill_type="solid")
    fill_zebrado = PatternFill(start_color=cor_linha_clara, end_color=cor_linha_clara, fill_type="solid")
    fill_totais = PatternFill(start_color="e2e8f0", end_color="e2e8f0", fill_type="solid")

    borda_fina = Side(style="thin", color=cor_borda)
    border_celula = Border(left=borda_fina, right=borda_fina, top=borda_fina, bottom=borda_fina)

    # 1. Cabeçalho Executivo
    ws.merge_cells("A1:N1")
    ws["A1"] = "REPÚBLICA DE ANGOLA — ADMINISTRAÇÃO GERAL TRIBUTÁRIA (AGT)"
    ws["A1"].font = fonte_titulo
    ws["A1"].fill = fill_titulo
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:N2")
    ws["A2"] = f"MAPA DE AQUISIÇÕES E FACTURAS DE FORNECEDORES — ANEXO MODELO 7 DO IVA | {empresa_nome} (NIF: {empresa_nif})"
    ws["A2"].font = fonte_sub
    ws["A2"].fill = fill_titulo
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    ws.append([])  # Linha 3 vazia

    # 2. Cabeçalhos da Tabela
    cabecalhos = [
        "N.º", "Tipo", "N.º Fatura", "Data Emissão", "Fornecedor / Emitente", "NIF Fornecedor",
        "Base Incidência (Kz)", "Taxa IVA (%)", "IVA Suportado (Kz)", "IVA Dedutível (Kz)",
        "Retenção 7% IVA", "Retenção 6.5% II", "Total Líquido (Kz)", "Conformidade AGT"
    ]
    ws.append(cabecalhos)
    ws.row_dimensions[4].height = 24

    for col_idx in range(1, len(cabecalhos) + 1):
        cell = ws.cell(row=4, column=col_idx)
        cell.font = fonte_cabecalho
        cell.fill = fill_cabecalho
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_celula

    # 3. Preenchimento de Linhas
    linha_atual = 5
    tot_base = Decimal('0.00')
    tot_iva = Decimal('0.00')
    tot_iva_ded = Decimal('0.00')
    tot_ret_iva = Decimal('0.00')
    tot_ret_ii = Decimal('0.00')
    tot_liq = Decimal('0.00')

    for idx, f in enumerate(qs, start=1):
        tot_base += f.base_incidencia or Decimal('0.00')
        tot_iva += f.valor_iva or Decimal('0.00')
        if f.iva_dedutivel:
            tot_iva_ded += f.valor_iva or Decimal('0.00')
        tot_ret_iva += f.retencao_iva or Decimal('0.00')
        tot_ret_ii += f.retencao_ii or Decimal('0.00')
        tot_liq += f.valor_liquido or Decimal('0.00')

        data_str = f.data_emissao.strftime('%d/%m/%Y') if f.data_emissao else "---"

        row_vals = [
            idx,
            f.tipo_documento,
            f.numero_factura,
            data_str,
            f.fornecedor_nome or "---",
            f.fornecedor_nif or "---",
            float(f.base_incidencia or 0),
            float(f.taxa_iva or 0),
            float(f.valor_iva or 0),
            float(f.valor_iva or 0) if f.iva_dedutivel else 0.0,
            float(f.retencao_iva or 0),
            float(f.retencao_ii or 0),
            float(f.valor_liquido or 0),
            f.get_conformidade_agt_display() if hasattr(f, 'get_conformidade_agt_display') else f.conformidade_agt
        ]
        ws.append(row_vals)
        ws.row_dimensions[linha_atual].height = 20

        for col_idx in range(1, len(row_vals) + 1):
            c = ws.cell(row=linha_atual, column=col_idx)
            c.font = fonte_dados
            c.border = border_celula
            if idx % 2 == 0:
                c.fill = fill_zebrado

            # Formatações numéricas
            if col_idx in [1, 2, 4, 14]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [3, 5, 6]:
                c.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx == 8:
                c.number_format = '0.00"%"'
                c.alignment = Alignment(horizontal="right", vertical="center")
            else:
                c.number_format = '#,##0.00'
                c.alignment = Alignment(horizontal="right", vertical="center")

        linha_atual += 1

    # Linha de Totais Finais
    totais_vals = [
        "TOTAIS", "", "", "", "", "",
        float(tot_base), "", float(tot_iva), float(tot_iva_ded),
        float(tot_ret_iva), float(tot_ret_ii), float(tot_liq), ""
    ]
    ws.append(totais_vals)
    ws.row_dimensions[linha_atual].height = 24
    for col_idx in range(1, len(totais_vals) + 1):
        c = ws.cell(row=linha_atual, column=col_idx)
        c.font = fonte_totais
        c.fill = fill_totais
        c.border = border_celula
        if col_idx in [7, 9, 10, 11, 12, 13]:
            c.number_format = '#,##0.00'
            c.alignment = Alignment(horizontal="right", vertical="center")
        else:
            c.alignment = Alignment(horizontal="center", vertical="center")

    # Ajuste de larguras de colunas
    larguras = {1: 6, 2: 8, 3: 18, 4: 13, 5: 32, 6: 15, 7: 19, 8: 12, 9: 17, 10: 17, 11: 17, 12: 17, 13: 19, 14: 16}
    for col_i, width in larguras.items():
        ws.column_dimensions[get_column_letter(col_i)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    excel_bytes = buffer.getvalue()
    buffer.close()

    filename = f"Mapa_Fornecedores_AGT_{empresa_nif}_{ano or 'Todos'}.xlsx"
    response = HttpResponse(
        excel_bytes,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response




